#!/usr/bin/env python3
"""Split a large XLSX into row chunks and scan chunks in parallel for MRD evidence."""

from __future__ import annotations

import argparse
import concurrent.futures
import csv
import hashlib
import json
import os
import re
from io import BytesIO
from pathlib import Path
from time import monotonic

import openpyxl

from mrdskill_extract_full_corpus_evidence import (
    CHUNK_COLUMNS,
    SOURCE_COLUMNS,
    fs_path,
    paper_id_from_name,
    safe_read_bytes,
    safe_size,
    scan_workbook_chunk,
    stable_id,
)
from mrdskill_extract_mrd_candidates import RAW_COLUMNS, build_row, load_candidate_table


EXTRA_COLUMNS = [
    "source_document_id",
    "source_document_type",
    "source_file",
    "source_path",
    "source_sheet",
    "source_row",
    "source_column",
    "source_parser",
    "source_parser_status",
]


def write_csv(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_jsonl(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


class Bar:
    def __init__(self, label: str, width: int = 28) -> None:
        self.label = label
        self.width = width
        self.total = 0
        self.done = 0
        self.candidates = 0
        self.started = monotonic()
        self.last = 0.0

    def add_total(self, n: int) -> None:
        self.total += n
        self.render()

    def advance(self, candidates: int = 0) -> None:
        self.done += 1
        self.candidates += candidates
        self.render()

    def render(self, final: bool = False) -> None:
        now = monotonic()
        if not final and now - self.last < 0.2:
            return
        self.last = now
        ratio = self.done / self.total if self.total else 0
        filled = min(self.width, int(ratio * self.width))
        bar = "#" * filled + "-" * (self.width - filled)
        elapsed = now - self.started
        rate = self.done / elapsed if elapsed else 0
        print(
            f"\r{self.label} [{bar}] chunks {self.done}/{self.total} "
            f"| candidates {self.candidates} | {rate:.2f} chunk/s",
            end="\n" if final else "",
            flush=True,
        )


def chunk_csv_path(chunk_dir: Path, source_document_id: str, sheet_name: str, start_row: int, end_row: int) -> Path:
    safe_sheet = re.sub(r"[^A-Za-z0-9_.-]+", "_", sheet_name)[:80] or "sheet"
    chunk_id = stable_id("CHUNK", source_document_id, sheet_name, start_row, end_row)
    return chunk_dir / f"{chunk_id}__{safe_sheet}__rows_{start_row}_{end_row}.csv"


def split_xlsx_to_chunks(
    xlsx_path: Path,
    chunk_dir: Path,
    source_document_id: str,
    chunk_rows: int,
    progress: Bar,
) -> list[dict[str, str]]:
    data = safe_read_bytes(xlsx_path)
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    chunk_logs: list[dict[str, str]] = []
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            chunk: list[tuple[int, tuple[object, ...]]] = []
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                chunk.append((row_idx, tuple(row)))
                if len(chunk) >= chunk_rows:
                    write_chunk(chunk_dir, xlsx_path, source_document_id, sheet_name, chunk, chunk_logs)
                    progress.add_total(1)
                    chunk = []
            if chunk:
                write_chunk(chunk_dir, xlsx_path, source_document_id, sheet_name, chunk, chunk_logs)
                progress.add_total(1)
    finally:
        workbook.close()
    return chunk_logs


def write_chunk(
    chunk_dir: Path,
    xlsx_path: Path,
    source_document_id: str,
    sheet_name: str,
    chunk: list[tuple[int, tuple[object, ...]]],
    chunk_logs: list[dict[str, str]],
) -> None:
    start_row = chunk[0][0]
    end_row = chunk[-1][0]
    path = chunk_csv_path(chunk_dir, source_document_id, sheet_name, start_row, end_row)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["__row_number__", "__sheet__", "__source_file__", "__cells__"])
        for row_idx, row in chunk:
            cells = ["" if value is None else str(value) for value in row]
            writer.writerow([row_idx, sheet_name, xlsx_path.name, json.dumps(cells, ensure_ascii=False)])
    chunk_logs.append(
        {
            "chunk_id": stable_id("CHUNK", source_document_id, sheet_name, start_row, end_row),
            "source_document_id": source_document_id,
            "paper_id": paper_id_from_name(xlsx_path.name),
            "source_file": xlsx_path.name,
            "source_sheet": sheet_name,
            "row_start": str(start_row),
            "row_end": str(end_row),
            "rows_seen": str(len(chunk)),
            "candidate_rows": "pending",
            "parser_status": "chunk_written",
            "note": str(path),
        }
    )


def scan_chunk_csv(path: Path) -> dict[str, object]:
    candidates = []
    rows_seen = 0
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            rows_seen += 1
            cells = json.loads(row["__cells__"])
            row_items = [(int(row["__row_number__"]), tuple(cells))]
            result = scan_workbook_chunk((row["__source_file__"], row["__sheet__"], int(row["__row_number__"]), row_items))
            candidates.extend(result["candidates"])  # type: ignore[arg-type]
    return {"path": str(path), "rows_seen": rows_seen, "candidate_rows": len(candidates), "candidates": candidates}


def main() -> int:
    parser = argparse.ArgumentParser(description="Split and scan a large XLSX for MRD evidence.")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--candidate-table", required=True)
    parser.add_argument("--chunk-dir", default="data/imported/excel_chunks")
    parser.add_argument("--out-csv", default="data/evidence/mrd_evidence_raw_excel.csv")
    parser.add_argument("--out-jsonl", default="data/evidence/mrd_evidence_raw_excel.jsonl")
    parser.add_argument("--sources-csv", default="data/imported/excel_source.csv")
    parser.add_argument("--chunk-log-csv", default="data/imported/excel_chunk_log.csv")
    parser.add_argument("--chunk-rows", type=int, default=20000)
    parser.add_argument("--workers", type=int, default=max(1, min(4, (os.cpu_count() or 2) - 1)))
    parser.add_argument("--max-evidence", type=int, default=0, help="0 means no cap.")
    parser.add_argument("--max-quote-chars", type=int, default=1200)
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx).resolve()
    chunk_dir = Path(args.chunk_dir).resolve() / xlsx_path.stem
    source_document_id = stable_id("SRC", str(xlsx_path))
    paper_id = paper_id_from_name(xlsx_path.name)
    candidate_rows = load_candidate_table(Path(args.candidate_table))
    metadata = candidate_rows.get(paper_id, {})

    split_bar = Bar("split")
    chunk_logs = split_xlsx_to_chunks(xlsx_path, chunk_dir, source_document_id, args.chunk_rows, split_bar)
    split_bar.render(final=True)

    chunk_paths = [Path(row["note"]) for row in chunk_logs]
    scan_bar = Bar("scan")
    scan_bar.add_total(len(chunk_paths))
    evidence_rows: list[dict[str, str]] = []
    chunk_log_by_path = {row["note"]: row for row in chunk_logs}

    with concurrent.futures.ProcessPoolExecutor(max_workers=max(1, args.workers)) as executor:
        futures = {executor.submit(scan_chunk_csv, path): path for path in chunk_paths}
        for future in concurrent.futures.as_completed(futures):
            result = future.result()
            path = str(result["path"])
            chunk_log_by_path[path]["candidate_rows"] = str(result["candidate_rows"])
            chunk_log_by_path[path]["parser_status"] = "scanned"
            for sheet, row_number, column_hits, snippet in result["candidates"]:  # type: ignore[union-attr]
                if args.max_evidence and len(evidence_rows) >= args.max_evidence:
                    continue
                row = build_row(
                    len(evidence_rows) + 1,
                    paper_id,
                    metadata,
                    "Supplementary_Table",
                    snippet,
                    args.max_quote_chars,
                )
                row["evidence_id"] = f"EV-XLSX-{len(evidence_rows) + 1:07d}"
                row.update(
                    {
                        "source_document_id": source_document_id,
                        "source_document_type": "supplementary_native_xlsx",
                        "source_file": xlsx_path.name,
                        "source_path": str(xlsx_path),
                        "source_sheet": sheet,
                        "source_row": row_number,
                        "source_column": column_hits,
                        "source_parser": "split_scan_excel",
                        "source_parser_status": "parsed",
                        "supplementary_file": xlsx_path.name,
                        "supplementary_table": sheet,
                    }
                )
                evidence_rows.append(row)
            scan_bar.advance(int(result["candidate_rows"]))
    scan_bar.render(final=True)

    source_row = {
        "source_document_id": source_document_id,
        "paper_id": paper_id,
        "source_document_type": "supplementary_native_xlsx",
        "source_file": xlsx_path.name,
        "source_path": str(xlsx_path),
        "extension": ".xlsx",
        "size_bytes": str(safe_size(xlsx_path)),
        "parser": "split_scan_excel",
        "parser_status": "parsed",
        "skip_reason": "not_reported",
        "evidence_count": str(len(evidence_rows)),
    }
    output_columns = RAW_COLUMNS + [column for column in EXTRA_COLUMNS if column not in RAW_COLUMNS]
    write_csv(Path(args.out_csv), evidence_rows, output_columns)
    write_jsonl(Path(args.out_jsonl), evidence_rows)
    write_csv(Path(args.sources_csv), [source_row], SOURCE_COLUMNS)
    write_csv(Path(args.chunk_log_csv), list(chunk_log_by_path.values()), CHUNK_COLUMNS)
    print(f"Wrote evidence: {Path(args.out_csv).resolve()} ({len(evidence_rows)} rows)")
    print(f"Wrote chunks: {Path(args.chunk_log_csv).resolve()} ({len(chunk_logs)} chunks)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
