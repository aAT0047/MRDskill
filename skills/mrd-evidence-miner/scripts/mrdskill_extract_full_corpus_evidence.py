#!/usr/bin/env python3
"""Extract candidate MRD evidence from all parseable kg_ready_no_pdf materials."""

from __future__ import annotations

import argparse
import concurrent.futures
import csv
import hashlib
import html
import json
import os
import re
import sys
import zipfile
from io import BytesIO
from pathlib import Path
from time import monotonic
from typing import Iterable

try:
    from mrdskill_extract_mrd_candidates import (
        ASSAY_PATTERNS,
        GENE_RE,
        KWRAW_RE,
        OUTCOME_TERMS,
        RAW_COLUMNS,
        VARIANT_RE,
        build_row,
        first_match,
        iter_mrd_snippets,
        load_candidate_table,
        load_converted_manifest,
        normalize_text,
    )
except ModuleNotFoundError:  # Backward compatibility for older unprefixed skill layouts.
    from extract_mrd_candidates import (
        ASSAY_PATTERNS,
        GENE_RE,
        KWRAW_RE,
        OUTCOME_TERMS,
        RAW_COLUMNS,
        VARIANT_RE,
        build_row,
        first_match,
        iter_mrd_snippets,
        load_candidate_table,
        load_converted_manifest,
        normalize_text,
    )

try:
    from bs4 import BeautifulSoup
except Exception:  # pragma: no cover - optional dependency
    BeautifulSoup = None

try:
    import openpyxl
except Exception:  # pragma: no cover - optional dependency
    openpyxl = None


EXTRA_COLUMNS = [
    "source_document_id",
    "source_document_type",
    "source_file",
    "source_path",
    "source_sheet",
    "source_row",
    "source_column",
    "source_parser",
    "source_parser_status",
]

SOURCE_COLUMNS = [
    "source_document_id",
    "paper_id",
    "source_document_type",
    "source_file",
    "source_path",
    "extension",
    "size_bytes",
    "parser",
    "parser_status",
    "skip_reason",
    "evidence_count",
]

CHUNK_COLUMNS = [
    "chunk_id",
    "source_document_id",
    "paper_id",
    "source_file",
    "source_sheet",
    "row_start",
    "row_end",
    "rows_seen",
    "candidate_rows",
    "parser_status",
    "note",
]

MRD_OR_OUTCOME_RE = re.compile(
    r"\bMRD\b|minimal residual disease|measurable residual disease|molecular residual disease|"
    r"\bctDNA\b|circulating tumou?r DNA|liquid biopsy|molecular relapse|recurrence|relapse|"
    r"progression-free survival|disease-free survival|overall survival|prognosis|clearance|persistence|"
    r"clonal hematopoiesis|clonal haematopoiesis|\bCHIP\b",
    re.IGNORECASE,
)

VARIANT_CONTEXT_RE = re.compile(
    r"\bvariant\b|\bmutation\b|\bmutant\b|\bVAF\b|variant allele|allele frequency|"
    r"\bAF\b|\bctDNA\b|\bcfDNA\b|\bplasma\b|\btumou?r\b|\btissue\b|\bsample\b|"
    r"baseline|postoperative|post[- ]?treatment|timepoint|MRD",
    re.IGNORECASE,
)


def stable_id(prefix: str, *parts: str) -> str:
    raw = "|".join(str(part) for part in parts if part is not None)
    digest = hashlib.sha1(raw.encode("utf-8", errors="ignore")).hexdigest()[:16]
    return f"{prefix}-{digest}"


def fs_path(path: Path) -> str:
    """Return a Windows long-path-safe string when needed."""
    path_str = str(path)
    if os.name != "nt":
        return path_str
    absolute = os.path.abspath(path_str)
    if absolute.startswith("\\\\?\\"):
        return absolute
    if absolute.startswith("\\\\"):
        return "\\\\?\\UNC\\" + absolute.lstrip("\\")
    return "\\\\?\\" + absolute


def safe_read_bytes(path: Path) -> bytes:
    with open(fs_path(path), "rb") as handle:
        return handle.read()


def safe_size(path: Path) -> int:
    try:
        return os.stat(fs_path(path)).st_size
    except OSError:
        return -1


def paper_id_from_name(name: str) -> str:
    match = KWRAW_RE.search(name)
    return match.group(1) if match else "not_reported"


def source_type_from_path(path: Path, root: Path) -> str:
    rel = str(path.relative_to(root)).replace("\\", "/").lower()
    suffix = path.suffix.lower().lstrip(".") or "no_ext"
    if rel.startswith("downloaded_pdfs/converted_from_pdf"):
        return "paper_converted_txt" if suffix == "txt" else f"paper_converted_{suffix}"
    if rel.startswith("downloaded_pdfs/publisher_fulltext_html_xml"):
        return f"publisher_fulltext_{suffix}"
    if rel.startswith("supplementary_materials/converted_from_pdf"):
        return "supplementary_converted_txt" if suffix == "txt" else f"supplementary_converted_{suffix}"
    if rel.startswith("supplementary_materials/native_non_pdf"):
        return f"supplementary_native_{suffix}"
    return f"unknown_{suffix}"


def text_from_html_or_xml(data: bytes, parser_hint: str) -> str:
    decoded = data.decode("utf-8", errors="replace")
    if BeautifulSoup is not None:
        parser = "xml" if parser_hint == "xml" else "lxml"
        soup = BeautifulSoup(decoded, parser)
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        return normalize_text(soup.get_text("\n"))
    decoded = re.sub(r"<(script|style).*?</\1>", " ", decoded, flags=re.IGNORECASE | re.DOTALL)
    decoded = re.sub(r"<[^>]+>", " ", decoded)
    return normalize_text(html.unescape(decoded))


def text_from_docx_bytes(data: bytes) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(BytesIO(data)) as archive:
        names = [
            name
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
            and ("document.xml" in name or "footnotes.xml" in name or "endnotes.xml" in name)
        ]
        for name in names:
            xml_text = archive.read(name).decode("utf-8", errors="replace")
            xml_text = re.sub(r"<w:tab\s*/>", "\t", xml_text)
            xml_text = re.sub(r"</w:p>", "\n", xml_text)
            xml_text = re.sub(r"<[^>]+>", " ", xml_text)
            parts.append(html.unescape(xml_text))
    return normalize_text("\n".join(parts))


def open_zip_from_path(path: Path) -> zipfile.ZipFile:
    return zipfile.ZipFile(fs_path(path))


def text_from_binary_strings(data: bytes) -> str:
    ascii_chunks = re.findall(rb"[\x20-\x7E]{4,}", data)
    utf16_text = data.decode("utf-16le", errors="ignore")
    utf16_chunks = re.findall(r"[\x20-\x7E]{4,}", utf16_text)
    combined = [chunk.decode("latin-1", errors="ignore") for chunk in ascii_chunks]
    combined.extend(utf16_chunks)
    return normalize_text("\n".join(combined))


def workbook_chunks(
    data: bytes,
    chunk_rows: int,
    max_rows_per_sheet: int,
    max_empty_rows: int,
) -> Iterable[tuple[str, int, list[tuple[int, tuple[object, ...]]], bool, str]]:
    if openpyxl is None:
        raise RuntimeError("openpyxl_not_available")
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            empty_rows = 0
            chunk: list[tuple[int, tuple[object, ...]]] = []
            truncated = False
            stop_reason = "completed"
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if max_rows_per_sheet and row_idx > max_rows_per_sheet:
                    truncated = True
                    stop_reason = f"max_rows_per_sheet:{max_rows_per_sheet}"
                    break
                if not any(value is not None and str(value).strip() for value in row):
                    empty_rows += 1
                    if max_empty_rows and empty_rows >= max_empty_rows:
                        truncated = True
                        stop_reason = f"max_empty_rows:{max_empty_rows}"
                        break
                    continue
                empty_rows = 0
                chunk.append((row_idx, tuple(row)))
                if len(chunk) >= chunk_rows:
                    yield sheet_name, chunk[0][0], chunk, False, "chunk"
                    chunk = []
            if chunk:
                yield sheet_name, chunk[0][0], chunk, truncated, stop_reason
            elif truncated:
                yield sheet_name, 0, [], truncated, stop_reason
    finally:
        workbook.close()


def scan_workbook_chunk(task: tuple[str, str, int, list[tuple[int, tuple[object, ...]]]]) -> dict[str, object]:
    source_label, sheet_name, first_row, row_items = task
    candidates: list[tuple[str, str, str, str]] = []
    rows_seen = 0
    for row_idx, row in row_items:
        rows_seen += 1
        cells = []
        hit_columns = []
        for col_idx, value in enumerate(row, start=1):
            if value is None:
                continue
            text = str(value).strip()
            if not text:
                continue
            cells.append(text)
            if is_candidate_text(text):
                hit_columns.append(str(col_idx))
        if not cells:
            continue
        row_text = normalize_text(" | ".join(cells))
        if is_candidate_text(row_text):
            candidates.append(
                (
                    sheet_name,
                    str(row_idx),
                    ";".join(hit_columns) if hit_columns else "not_reported",
                    f"{source_label} | {row_text}",
                )
            )
    last_row = row_items[-1][0] if row_items else first_row
    return {
        "sheet": sheet_name,
        "row_start": str(first_row),
        "row_end": str(last_row),
        "rows_seen": rows_seen,
        "candidate_rows": len(candidates),
        "candidates": candidates,
    }


def rows_from_workbook_bytes(
    data: bytes,
    source_label: str,
    max_rows_per_sheet: int,
    max_empty_rows: int,
    chunk_rows: int,
    workers: int,
    progress: "Progress",
    source_document_id: str,
    paper_id: str,
    chunk_log_rows: list[dict[str, str]],
) -> list[tuple[str, str, str, str]]:
    candidates: list[tuple[str, str, str, str]] = []
    max_workers = max(1, workers)

    def handle_result(result: dict[str, object], note_tuple: tuple[str, int, int, bool, str]) -> None:
        candidates.extend(result["candidates"])  # type: ignore[arg-type]
        add_chunk_log(
            chunk_log_rows,
            source_document_id,
            paper_id,
            source_label,
            result,
            [note_tuple],
        )
        progress.advance_excel_chunk()

    if max_workers == 1:
        for sheet_name, first_row, chunk, truncated, note in workbook_chunks(
            data,
            chunk_rows,
            max_rows_per_sheet,
            max_empty_rows,
        ):
            row_end = chunk[-1][0] if chunk else first_row
            if not chunk:
                add_empty_chunk_log(chunk_log_rows, source_document_id, paper_id, source_label, sheet_name, first_row, row_end, note)
                continue
            progress.add_excel_chunks(1)
            result = scan_workbook_chunk((source_label, sheet_name, first_row, chunk))
            handle_result(result, (sheet_name, first_row, row_end, truncated, note))
    else:
        with concurrent.futures.ProcessPoolExecutor(max_workers=max_workers) as executor:
            pending: dict[concurrent.futures.Future[dict[str, object]], tuple[str, int, int, bool, str]] = {}

            def drain_completed(block: bool = False) -> None:
                if not pending:
                    return
                timeout = None if block else 0
                done, _ = concurrent.futures.wait(
                    pending,
                    timeout=timeout,
                    return_when=concurrent.futures.FIRST_COMPLETED,
                )
                for future in done:
                    note_tuple = pending.pop(future)
                    handle_result(future.result(), note_tuple)

            for sheet_name, first_row, chunk, truncated, note in workbook_chunks(
                data,
                chunk_rows,
                max_rows_per_sheet,
                max_empty_rows,
            ):
                row_end = chunk[-1][0] if chunk else first_row
                if not chunk:
                    add_empty_chunk_log(chunk_log_rows, source_document_id, paper_id, source_label, sheet_name, first_row, row_end, note)
                    continue
                progress.add_excel_chunks(1)
                future = executor.submit(scan_workbook_chunk, (source_label, sheet_name, first_row, chunk))
                pending[future] = (sheet_name, first_row, row_end, truncated, note)
                while len(pending) >= max_workers * 3:
                    drain_completed(block=True)
                drain_completed(block=False)

            while pending:
                drain_completed(block=True)
    return candidates


def add_empty_chunk_log(
    chunk_log_rows: list[dict[str, str]],
    source_document_id: str,
    paper_id: str,
    source_label: str,
    sheet_name: str,
    first_row: int,
    row_end: int,
    note: str,
) -> None:
    chunk_log_rows.append(
        {
            "chunk_id": stable_id("CHUNK", source_document_id, sheet_name, first_row, row_end),
            "source_document_id": source_document_id,
            "paper_id": paper_id,
            "source_file": source_label,
            "source_sheet": sheet_name,
            "row_start": str(first_row),
            "row_end": str(row_end),
            "rows_seen": "0",
            "candidate_rows": "0",
            "parser_status": "truncated",
            "note": note,
        }
    )


def add_chunk_log(
    chunk_log_rows: list[dict[str, str]],
    source_document_id: str,
    paper_id: str,
    source_label: str,
    result: dict[str, object],
    chunk_notes: list[tuple[str, int, int, bool, str]],
) -> None:
    sheet = str(result["sheet"])
    row_start = str(result["row_start"])
    row_end = str(result["row_end"])
    note = "completed"
    parser_status = "parsed"
    for note_sheet, note_start, note_end, truncated, note_text in chunk_notes:
        if note_sheet == sheet and str(note_start) == row_start and str(note_end) == row_end:
            note = note_text
            parser_status = "truncated" if truncated else "parsed"
            break
    chunk_log_rows.append(
        {
            "chunk_id": stable_id("CHUNK", source_document_id, sheet, row_start, row_end),
            "source_document_id": source_document_id,
            "paper_id": paper_id,
            "source_file": source_label,
            "source_sheet": sheet,
            "row_start": row_start,
            "row_end": row_end,
            "rows_seen": str(result["rows_seen"]),
            "candidate_rows": str(result["candidate_rows"]),
            "parser_status": parser_status,
            "note": note,
        }
    )


def rows_from_delimited_text(
    text: str,
    source_label: str,
    delimiter: str | None,
) -> Iterable[tuple[str, str, str, str]]:
    sample = text[:4096]
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff(sample)
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(text.splitlines(), delimiter=delimiter)
    for row_idx, row in enumerate(reader, start=1):
        cells = [cell.strip() for cell in row if cell and cell.strip()]
        if not cells:
            continue
        row_text = normalize_text(" | ".join(cells))
        if is_candidate_text(row_text):
            yield "table", str(row_idx), "not_reported", f"{source_label} | {row_text}"


class Progress:
    def __init__(self, total_files: int, enabled: bool = True, width: int = 28) -> None:
        self.total_files = total_files
        self.enabled = enabled
        self.width = width
        self.files_done = 0
        self.evidence_rows = 0
        self.excel_chunks_total = 0
        self.excel_chunks_done = 0
        self.started = monotonic()
        self.last_rendered = 0.0
        self.current = ""

    def add_excel_chunks(self, count: int) -> None:
        self.excel_chunks_total += count
        self.render()

    def advance_excel_chunk(self) -> None:
        self.excel_chunks_done += 1
        self.render()

    def advance_file(self, evidence_count: int, current: str) -> None:
        self.files_done += 1
        self.evidence_rows += evidence_count
        self.current = current
        self.render()

    def render(self, force_newline: bool = False) -> None:
        if not self.enabled:
            return
        now = monotonic()
        if not force_newline and now - self.last_rendered < 0.2:
            return
        self.last_rendered = now
        ratio = self.files_done / self.total_files if self.total_files else 1
        filled = min(self.width, int(ratio * self.width))
        bar = "#" * filled + "-" * (self.width - filled)
        elapsed = now - self.started
        rate = self.files_done / elapsed if elapsed else 0
        msg = (
            f"\r[{bar}] {self.files_done}/{self.total_files} files "
            f"| evidence {self.evidence_rows} "
            f"| xlsx chunks {self.excel_chunks_done}/{self.excel_chunks_total} "
            f"| {rate:.2f} file/s "
            f"| {self.current[:58]}"
        )
        print(msg, end="\n" if force_newline else "", flush=True)


def is_candidate_text(text: str) -> bool:
    if MRD_OR_OUTCOME_RE.search(text):
        return True
    if GENE_RE.search(text) and VARIANT_RE.search(text):
        return True
    if GENE_RE.search(text) and VARIANT_CONTEXT_RE.search(text):
        return True
    if VARIANT_RE.search(text) and VARIANT_CONTEXT_RE.search(text):
        return True
    if GENE_RE.search(text) and (VARIANT_RE.search(text) or any(term in text.lower() for term in OUTCOME_TERMS)):
        return True
    if first_match(ASSAY_PATTERNS, text) != "not_reported" and (
        VARIANT_RE.search(text) or GENE_RE.search(text) or MRD_OR_OUTCOME_RE.search(text)
    ):
        return True
    return False


def emit_rows_from_text(
    text: str,
    row_start: int,
    paper_id: str,
    metadata: dict[str, str],
    source_info: dict[str, str],
    max_snippets: int,
    context_chars: int,
    max_quote_chars: int,
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for section, snippet in iter_mrd_snippets(text, context_chars):
        key = hashlib.sha1(snippet[:1000].encode("utf-8", errors="ignore")).hexdigest()
        if key in seen:
            continue
        seen.add(key)
        row = build_row(row_start + len(rows), paper_id, metadata, section, snippet, max_quote_chars)
        row.update(source_info)
        rows.append(row)
        if len(rows) >= max_snippets:
            break
    return rows


def emit_rows_from_table_candidates(
    candidates: Iterable[tuple[str, str, str, str]],
    row_start: int,
    paper_id: str,
    metadata: dict[str, str],
    source_info: dict[str, str],
    max_snippets: int,
    max_quote_chars: int,
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for sheet, row_number, column_hits, snippet in candidates:
        key = hashlib.sha1(snippet[:1000].encode("utf-8", errors="ignore")).hexdigest()
        if key in seen:
            continue
        seen.add(key)
        row = build_row(row_start + len(rows), paper_id, metadata, "Supplementary_Table", snippet, max_quote_chars)
        row.update(source_info)
        row["source_sheet"] = sheet
        row["source_row"] = row_number
        row["source_column"] = column_hits or "not_reported"
        row["supplementary_file"] = source_info["source_file"]
        row["supplementary_table"] = sheet
        rows.append(row)
        if len(rows) >= max_snippets:
            break
    return rows


def source_record(
    source_document_id: str,
    paper_id: str,
    source_document_type: str,
    source_file: str,
    source_path: str,
    extension: str,
    size_bytes: int,
    parser: str,
    parser_status: str,
    skip_reason: str,
    evidence_count: int = 0,
) -> dict[str, str]:
    return {
        "source_document_id": source_document_id,
        "paper_id": paper_id,
        "source_document_type": source_document_type,
        "source_file": source_file,
        "source_path": source_path,
        "extension": extension,
        "size_bytes": str(size_bytes),
        "parser": parser,
        "parser_status": parser_status,
        "skip_reason": skip_reason,
        "evidence_count": str(evidence_count),
    }


def process_file(
    path: Path,
    root: Path,
    candidate_rows: dict[str, dict[str, str]],
    row_start: int,
    args: argparse.Namespace,
    progress: Progress,
    chunk_log_rows: list[dict[str, str]],
) -> tuple[list[dict[str, str]], dict[str, str]]:
    paper_id = paper_id_from_name(path.name)
    metadata = candidate_rows.get(paper_id, {})
    source_document_type = source_type_from_path(path, root)
    source_document_id = stable_id("SRC", str(path.relative_to(root)))
    ext = path.suffix.lower()
    source_info = {
        "source_document_id": source_document_id,
        "source_document_type": source_document_type,
        "source_file": path.name,
        "source_path": str(path),
        "source_sheet": "not_reported",
        "source_row": "not_reported",
        "source_column": "not_reported",
        "source_parser": "not_reported",
        "source_parser_status": "not_reported",
    }
    rows: list[dict[str, str]] = []
    parser = "not_reported"
    status = "parsed"
    reason = "not_reported"

    try:
        data = safe_read_bytes(path)
        if ext in {".txt"}:
            parser = "text"
            text = normalize_text(data.decode("utf-8", errors="replace"))
            source_info.update({"source_parser": parser, "source_parser_status": status})
            rows = emit_rows_from_text(
                text,
                row_start,
                paper_id,
                metadata,
                source_info,
                args.max_snippets_per_document,
                args.context_chars,
                args.max_quote_chars,
            )
        elif ext in {".html", ".htm", ".xml"}:
            parser = "beautifulsoup"
            text = text_from_html_or_xml(data, "xml" if ext == ".xml" else "html")
            source_info.update({"source_parser": parser, "source_parser_status": status})
            rows = emit_rows_from_text(
                text,
                row_start,
                paper_id,
                metadata,
                source_info,
                args.max_snippets_per_document,
                args.context_chars,
                args.max_quote_chars,
            )
        elif ext in {".csv", ".tsv"}:
            parser = "csv"
            text = data.decode("utf-8-sig", errors="replace")
            delimiter = "\t" if ext == ".tsv" else None
            source_info.update({"source_parser": parser, "source_parser_status": status})
            candidates = rows_from_delimited_text(text, path.name, delimiter)
            rows = emit_rows_from_table_candidates(
                candidates,
                row_start,
                paper_id,
                metadata,
                source_info,
                args.max_table_evidence_per_document,
                args.max_quote_chars,
            )
        elif ext == ".xlsx":
            parser = "openpyxl"
            source_info.update({"source_parser": parser, "source_parser_status": status})
            candidates = rows_from_workbook_bytes(
                data,
                path.name,
                args.max_rows_per_sheet,
                args.max_empty_rows,
                args.excel_chunk_rows,
                args.excel_workers,
                progress,
                source_document_id,
                paper_id,
                chunk_log_rows,
            )
            rows = emit_rows_from_table_candidates(
                candidates,
                row_start,
                paper_id,
                metadata,
                source_info,
                args.max_table_evidence_per_document,
                args.max_quote_chars,
            )
        elif ext == ".docx":
            parser = "docx_zip_xml"
            text = text_from_docx_bytes(data)
            source_info.update({"source_parser": parser, "source_parser_status": status})
            rows = emit_rows_from_text(
                text,
                row_start,
                paper_id,
                metadata,
                source_info,
                args.max_snippets_per_document,
                args.context_chars,
                args.max_quote_chars,
            )
        elif ext in {".doc", ".xls"}:
            parser = "binary_strings"
            status = "parsed_best_effort"
            text = text_from_binary_strings(data)
            source_info.update({"source_parser": parser, "source_parser_status": status})
            rows = emit_rows_from_text(
                text,
                row_start,
                paper_id,
                metadata,
                source_info,
                args.max_snippets_per_document,
                args.context_chars,
                args.max_quote_chars,
            )
        elif ext == ".zip":
            parser = "zip"
            source_info.update({"source_parser": parser, "source_parser_status": status})
            rows = process_zip_entries(
                path,
                root,
                candidate_rows,
                row_start,
                args,
                source_document_id,
                progress,
                chunk_log_rows,
            )
        else:
            status = "skipped"
            reason = f"unsupported_extension:{ext or '<no_ext>'}"
    except Exception as exc:
        status = "error"
        reason = f"{type(exc).__name__}: {exc}"

    for offset, row in enumerate(rows, start=row_start):
        row["evidence_id"] = f"EV-FULL-{offset:07d}"

    source_info["source_parser_status"] = status
    record = source_record(
        source_document_id,
        paper_id,
        source_document_type,
        path.name,
        str(path),
        ext or "<no_ext>",
        safe_size(path),
        parser,
        status,
        reason,
        len(rows),
    )
    return rows, record


def process_zip_entries(
    path: Path,
    root: Path,
    candidate_rows: dict[str, dict[str, str]],
    row_start: int,
    args: argparse.Namespace,
    parent_source_document_id: str,
    progress: Progress,
    chunk_log_rows: list[dict[str, str]],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    paper_id = paper_id_from_name(path.name)
    metadata = candidate_rows.get(paper_id, {})
    with open_zip_from_path(path) as archive:
        for entry in archive.infolist():
            if entry.is_dir():
                continue
            if len(rows) >= args.max_table_evidence_per_document:
                break
            ext = Path(entry.filename).suffix.lower()
            if ext not in {".txt", ".csv", ".tsv", ".html", ".htm", ".xml", ".docx", ".xlsx", ".doc", ".xls"}:
                continue
            data = archive.read(entry)
            source_document_id = stable_id("SRC", str(path.relative_to(root)), entry.filename)
            source_info = {
                "source_document_id": source_document_id,
                "source_document_type": f"zip_entry_{ext.lstrip('.')}",
                "source_file": f"{path.name}!{entry.filename}",
                "source_path": f"{path}!{entry.filename}",
                "source_sheet": "not_reported",
                "source_row": "not_reported",
                "source_column": "not_reported",
                "source_parser": "zip_entry",
                "source_parser_status": f"parent:{parent_source_document_id}",
            }
            try:
                if ext in {".txt", ".doc", ".xls"}:
                    text = data.decode("utf-8", errors="replace") if ext == ".txt" else text_from_binary_strings(data)
                    rows.extend(
                        emit_rows_from_text(
                            text,
                            row_start + len(rows),
                            paper_id,
                            metadata,
                            source_info,
                            max(0, args.max_table_evidence_per_document - len(rows)),
                            args.context_chars,
                            args.max_quote_chars,
                        )
                    )
                elif ext in {".html", ".htm", ".xml"}:
                    text = text_from_html_or_xml(data, "xml" if ext == ".xml" else "html")
                    rows.extend(
                        emit_rows_from_text(
                            text,
                            row_start + len(rows),
                            paper_id,
                            metadata,
                            source_info,
                            max(0, args.max_table_evidence_per_document - len(rows)),
                            args.context_chars,
                            args.max_quote_chars,
                        )
                    )
                elif ext == ".docx":
                    text = text_from_docx_bytes(data)
                    rows.extend(
                        emit_rows_from_text(
                            text,
                            row_start + len(rows),
                            paper_id,
                            metadata,
                            source_info,
                            max(0, args.max_table_evidence_per_document - len(rows)),
                            args.context_chars,
                            args.max_quote_chars,
                        )
                    )
                elif ext in {".csv", ".tsv"}:
                    text = data.decode("utf-8-sig", errors="replace")
                    candidates = rows_from_delimited_text(text, entry.filename, "\t" if ext == ".tsv" else None)
                    rows.extend(
                        emit_rows_from_table_candidates(
                            candidates,
                            row_start + len(rows),
                            paper_id,
                            metadata,
                            source_info,
                            max(0, args.max_table_evidence_per_document - len(rows)),
                            args.max_quote_chars,
                        )
                    )
                elif ext == ".xlsx":
                    candidates = rows_from_workbook_bytes(
                        data,
                        entry.filename,
                        args.max_rows_per_sheet,
                        args.max_empty_rows,
                        args.excel_chunk_rows,
                        args.excel_workers,
                        progress,
                        source_document_id,
                        paper_id,
                        chunk_log_rows,
                    )
                    rows.extend(
                        emit_rows_from_table_candidates(
                            candidates,
                            row_start + len(rows),
                            paper_id,
                            metadata,
                            source_info,
                            max(0, args.max_table_evidence_per_document - len(rows)),
                            args.max_quote_chars,
                        )
                    )
            except Exception:
                continue
    return rows


def iter_all_source_files(root: Path) -> Iterable[Path]:
    target_dirs = [
        root / "downloaded_pdfs" / "converted_from_pdf",
        root / "downloaded_pdfs" / "publisher_fulltext_html_xml",
        root / "supplementary_materials" / "converted_from_pdf",
        root / "supplementary_materials" / "native_non_pdf",
    ]
    for directory in target_dirs:
        if not directory.exists():
            continue
        for dirpath, _, filenames in os.walk(directory):
            for filename in filenames:
                path = Path(dirpath) / filename
                suffix = path.suffix.lower()
                if suffix in {".json"}:
                    continue
                if suffix in {".html", ".htm", ".xml", ".txt", ".csv", ".tsv", ".xlsx", ".docx", ".doc", ".xls", ".zip"}:
                    yield path


def write_csv(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_jsonl(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract MRD evidence from every parseable kg_ready_no_pdf source.")
    parser.add_argument("--corpus-dir", required=True)
    parser.add_argument("--out-csv", default="data/evidence/mrd_evidence_raw_full.csv")
    parser.add_argument("--out-jsonl", default="data/evidence/mrd_evidence_raw_full.jsonl")
    parser.add_argument("--sources-csv", default="data/imported/full_corpus_sources.csv")
    parser.add_argument("--chunk-log-csv", default="data/imported/full_corpus_excel_chunks.csv")
    parser.add_argument("--report-json", default="data/imported/full_corpus_ingest_report.json")
    parser.add_argument("--max-files", type=int, default=0, help="0 means all files.")
    parser.add_argument("--only-file-contains", default="", help="Process only files whose name contains this text.")
    parser.add_argument("--max-snippets-per-document", type=int, default=6)
    parser.add_argument("--max-table-evidence-per-document", type=int, default=200)
    parser.add_argument(
        "--max-rows-per-sheet",
        type=int,
        default=0,
        help="Rows scanned per workbook sheet. 0 means all rows.",
    )
    parser.add_argument("--max-empty-rows", type=int, default=0, help="Stop a sheet after this many empty rows. 0 disables early stop.")
    parser.add_argument("--excel-chunk-rows", type=int, default=20000)
    parser.add_argument("--excel-workers", type=int, default=max(1, min(4, (os.cpu_count() or 2) - 1)))
    parser.add_argument("--no-progress", action="store_true")
    parser.add_argument("--context-chars", type=int, default=900)
    parser.add_argument("--max-quote-chars", type=int, default=1200)
    parser.add_argument("--progress-every", type=int, default=250)
    args = parser.parse_args()

    root = Path(args.corpus_dir).resolve()
    if not root.exists():
        raise SystemExit(f"Corpus directory not found: {root}")
    manifests = root / "manifests"
    candidate_rows = load_candidate_table(manifests / "keyword_research_candidate_table.csv")

    # Keep this load as a sanity check for converted main-paper layout.
    load_converted_manifest(manifests / "downloaded_pdfs_converted_manifest.csv", root)

    evidence_rows: list[dict[str, str]] = []
    source_rows: list[dict[str, str]] = []
    chunk_log_rows: list[dict[str, str]] = []
    status_counts: dict[str, int] = {}
    source_type_counts: dict[str, int] = {}

    files = list(iter_all_source_files(root))
    if args.only_file_contains:
        files = [path for path in files if args.only_file_contains.lower() in path.name.lower()]
    if args.max_files:
        files = files[: args.max_files]

    progress = Progress(len(files), enabled=not args.no_progress)
    for index, path in enumerate(files, start=1):
        rows, source = process_file(
            path,
            root,
            candidate_rows,
            len(evidence_rows) + 1,
            args,
            progress,
            chunk_log_rows,
        )
        evidence_rows.extend(rows)
        source_rows.append(source)
        status_counts[source["parser_status"]] = status_counts.get(source["parser_status"], 0) + 1
        source_type_counts[source["source_document_type"]] = source_type_counts.get(source["source_document_type"], 0) + 1
        progress.advance_file(len(rows), path.name)
        if args.progress_every and index % args.progress_every == 0:
            progress.render(force_newline=True)

    progress.render(force_newline=True)
    output_columns = RAW_COLUMNS + [column for column in EXTRA_COLUMNS if column not in RAW_COLUMNS]
    write_csv(Path(args.out_csv), evidence_rows, output_columns)
    write_jsonl(Path(args.out_jsonl), evidence_rows)
    write_csv(Path(args.sources_csv), source_rows, SOURCE_COLUMNS)
    write_csv(Path(args.chunk_log_csv), chunk_log_rows, CHUNK_COLUMNS)
    report = {
        "corpus_dir": str(root),
        "source_files_seen": len(files),
        "evidence_rows": len(evidence_rows),
        "excel_chunk_rows": len(chunk_log_rows),
        "excel_chunk_size_rows": args.excel_chunk_rows,
        "excel_workers": args.excel_workers,
        "max_rows_per_sheet": args.max_rows_per_sheet,
        "max_table_evidence_per_document": args.max_table_evidence_per_document,
        "parser_status_counts": status_counts,
        "source_document_type_counts": source_type_counts,
        "outputs": {
            "raw_csv": str(Path(args.out_csv).resolve()),
            "raw_jsonl": str(Path(args.out_jsonl).resolve()),
            "sources_csv": str(Path(args.sources_csv).resolve()),
            "chunk_log_csv": str(Path(args.chunk_log_csv).resolve()),
        },
    }
    Path(args.report_json).parent.mkdir(parents=True, exist_ok=True)
    Path(args.report_json).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote full raw evidence CSV: {Path(args.out_csv).resolve()} ({len(evidence_rows)} rows)")
    print(f"Wrote source document CSV: {Path(args.sources_csv).resolve()} ({len(source_rows)} rows)")
    print(f"Wrote ingest report: {Path(args.report_json).resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
