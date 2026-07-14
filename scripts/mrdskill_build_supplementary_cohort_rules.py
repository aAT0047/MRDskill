#!/usr/bin/env python3
"""Build a supplementary-material cohort layer and candidate gene rules.

This script mines already-downloaded supplementary tables referenced by the
local source manifest. It is intentionally conservative: rows are kept only
when a table has sample-level identifiers plus gene or genomic variant fields.
The output is a cohort/statistical layer beside the literature-evidence KG.
"""

from __future__ import annotations

import argparse
import bisect
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from scipy.stats import fisher_exact

from mrdskill_project_paths import artifact_path


MISSING = {"", ".", "na", "n/a", "none", "null", "unknown", "uncertain", "unavailable", "not_reported", "nan"}
GENE_RE = re.compile(r"^[A-Z][A-Z0-9.-]{1,15}$")
COORD_RE = re.compile(r"^(?:chr)?([0-9]{1,2}|X|Y|M|MT)[:_\-\s]+([0-9][0-9,]*)", re.I)
MAX_RULE_SUPPORT_EVENTS = 80
MAX_RULE_SUPPORT_SOURCES = 100
GENE_TOKEN_BLACKLIST = {
    "ALTERATION", "ALTERATIONS", "AMPLIFICATION", "AMPLIFIED", "BLOOD", "CTDNA", "CTDNA-STATUS",
    "DELETION", "DELETED", "FUSION", "GENE", "GENES", "HIGH", "INDEL", "INSERTION", "LNCRNA",
    "LOW", "MEDIUM", "MRD", "MSI", "MSI-H", "MSI-L", "MSS", "MUT", "MUTANT", "MUTATED",
    "MUTATION", "MUTATIONS", "NEGATIVE", "NO", "NONE", "NORMAL", "OTHER", "PLASMA", "POSITIVE",
    "PROMOTER", "PROTEIN", "SCORE", "SERUM", "SILENT", "SNV", "STATUS", "TISSUE", "TMB",
    "TUMOR", "UNKNOWN", "VARIANT", "VARIANTS", "WILD", "WILDTYPE", "WT", "YES",
}


def add_limited(target: set[str], values: Any, limit: int) -> None:
    if len(target) >= limit:
        return
    if isinstance(values, str):
        iterable = [item for item in values.split(";") if item]
    else:
        iterable = values or []
    for value in iterable:
        text = clean(value)
        if text:
            target.add(text)
        if len(target) >= limit:
            break


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Mine supplementary sample-gene matrices and association rules.")
    parser.add_argument("--sources-csv", default="data/imported/full_corpus_sources.csv")
    parser.add_argument("--excel-chunks-csv", default="data/imported/full_corpus_excel_chunks.csv")
    parser.add_argument("--evidence-csv", default="data/evidence/mrd_evidence_scored_full.csv")
    parser.add_argument("--out-variant-csv", default="data/cohort/supplementary_sample_variant_matrix.csv")
    parser.add_argument("--out-gene-csv", default="data/cohort/supplementary_sample_gene_matrix.csv")
    parser.add_argument("--out-feature-csv", default="data/cohort/supplementary_sample_feature_matrix.csv")
    parser.add_argument("--out-baseline-csv", default="data/qc_rules/literature_cancer_gene_baselines.csv")
    parser.add_argument("--out-cooccur-csv", default="data/qc_rules/literature_gene_cooccurrence_rules.csv")
    parser.add_argument("--out-mutex-csv", default="data/qc_rules/literature_gene_mutex_rules.csv")
    parser.add_argument("--out-report-md", default="docs/qc/supplementary_cohort_rule_build_report.md")
    parser.add_argument("--from-feature-csv", default="", help="Resume rule generation from an existing supplementary_sample_feature_matrix.csv.")
    parser.add_argument("--from-variant-csv", default="", help="Rebuild sample gene/feature matrices and rules from an existing supplementary_sample_variant_matrix.csv.")
    parser.add_argument("--ncbi-grch37-json", default="data/annotation/ncbi/GCF_000001405.25_GRCh37.p13_feature_table.GRCh37.gene_intervals.json")
    parser.add_argument("--ncbi-grch38-json", default="data/annotation/ncbi/GCF_000001405.40_GRCh38.p14_feature_table.GRCh38.gene_intervals.json")
    parser.add_argument("--coordinate-assemblies", default="GRCh37,GRCh38", help="Comma-separated NCBI assemblies used for coordinate-only gene rescue.")
    parser.add_argument("--max-files", type=int, default=0, help="Debug limit; 0 means no limit.")
    parser.add_argument("--max-excel-sources", type=int, default=200, help="Limit Excel sources for practical first-pass builds; 0 means no limit.")
    parser.add_argument("--max-sheets-per-source", type=int, default=12, help="Limit candidate sheets per workbook; 0 means no limit.")
    parser.add_argument("--max-candidate-sheets-total", type=int, default=300, help="Global limit for Excel candidate sheets; 0 means no limit.")
    parser.add_argument("--max-source-size-mb", type=float, default=25.0, help="Skip larger supplementary files in first-pass builds; 0 means no limit.")
    parser.add_argument("--extensions", default=".xlsx,.xls,.csv,.tsv", help="Comma-separated source extensions to scan.")
    parser.add_argument("--candidate-sheets-only", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--min-candidate-rows", type=int, default=1)
    parser.add_argument("--max-rows-per-sheet", type=int, default=200000)
    parser.add_argument("--min-samples", type=int, default=20)
    parser.add_argument("--min-pair-support", type=int, default=3)
    parser.add_argument("--min-lift", type=float, default=1.5)
    parser.add_argument("--max-mutex-odds-ratio", type=float, default=0.35)
    parser.add_argument("--mutex-max-p-value", type=float, default=0.05)
    parser.add_argument("--min-mutex-feature-count", type=int, default=5)
    parser.add_argument("--min-mutex-expected", type=float, default=2.0)
    parser.add_argument("--mutex-feature-types", default="gene", help="Comma-separated feature types eligible for mutual-exclusion rules.")
    parser.add_argument("--max-features-per-sample-for-pairs", type=int, default=100)
    parser.add_argument("--skip-not-reported-pair-rules", action=argparse.BooleanOptionalAction, default=True)
    return parser.parse_args()


def load_ncbi_index(paths_by_assembly: dict[str, Path]) -> dict[str, dict[str, dict[str, Any]]]:
    indexes: dict[str, dict[str, dict[str, Any]]] = {}
    for assembly, path in paths_by_assembly.items():
        if not path.exists():
            continue
        raw = json.loads(path.read_text(encoding="utf-8"))
        chrom_index: dict[str, dict[str, Any]] = {}
        for chrom, intervals in raw.items():
            rows = sorted(intervals, key=lambda item: int(item.get("start") or 0))
            chrom_index[normalize_chrom(chrom)] = {"starts": [int(item.get("start") or 0) for item in rows], "rows": rows}
        indexes[assembly] = chrom_index
    return indexes


def ncbi_gene_symbols_from_index(ncbi_index: dict[str, dict[str, dict[str, Any]]]) -> set[str]:
    symbols: set[str] = set()
    for assembly in ncbi_index.values():
        for chrom_index in assembly.values():
            for row in chrom_index.get("rows", []):
                symbol = norm_gene(row.get("symbol"))
                if symbol:
                    symbols.add(symbol)
    return symbols


def normalize_chrom(value: Any) -> str:
    text = clean(value).upper()
    text = re.sub(r"^CHR", "", text, flags=re.I)
    if text in {"M", "MT", "MITO"}:
        return "MT"
    return text


def annotate_coordinate_gene(chrom: str, pos: str, ncbi_index: dict[str, dict[str, dict[str, Any]]], assemblies: list[str]) -> dict[str, str]:
    chrom_norm = normalize_chrom(chrom)
    try:
        position = int(clean(pos).replace(",", ""))
    except ValueError:
        return {"coordinate_gene": "", "coordinate_gene_assembly": "", "coordinate_gene_status": "not_annotatable"}
    if not chrom_norm or not position:
        return {"coordinate_gene": "", "coordinate_gene_assembly": "", "coordinate_gene_status": "not_annotatable"}
    for assembly in assemblies:
        chrom_index = ncbi_index.get(assembly, {}).get(chrom_norm)
        if not chrom_index:
            continue
        starts = chrom_index["starts"]
        rows = chrom_index["rows"]
        idx = bisect.bisect_right(starts, position)
        hits = []
        scan_start = max(0, idx - 30)
        for row in rows[scan_start : min(len(rows), idx + 30)]:
            start = int(row.get("start") or 0)
            end = int(row.get("end") or 0)
            if start <= position <= end:
                hits.append(row)
        if hits:
            hits.sort(key=lambda item: (item.get("gene_type") != "protein_coding", len(str(item.get("symbol") or ""))))
            symbols = sorted({clean(item.get("symbol")) for item in hits if clean(item.get("symbol"))})
            return {
                "coordinate_gene": ";".join(symbols),
                "coordinate_gene_assembly": assembly,
                "coordinate_gene_status": "ncbi_interval_match",
            }
    return {"coordinate_gene": "", "coordinate_gene_assembly": ";".join(assemblies), "coordinate_gene_status": "no_gene_interval_match"}


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def norm_missing(value: Any) -> bool:
    return clean(value).lower() in MISSING


def norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", clean(value).lower()).strip("_")


def norm_gene(value: Any) -> str:
    text = clean(value).upper()
    if norm_missing(text):
        return ""
    text = re.sub(r"\s+", "", text)
    text = text.split("|")[0] if "|" in text else text
    if text in GENE_TOKEN_BLACKLIST:
        return ""
    return text if GENE_RE.match(text) else ""


def split_gene_values(value: Any) -> list[str]:
    text = clean(value)
    if norm_missing(text):
        return []
    genes: set[str] = set()
    for token in re.split(r"[,;/\s]+", text):
        gene = norm_gene(token)
        if gene and not gene.startswith(("CHR", "EXON", "RS")):
            genes.add(gene)
    return sorted(genes)


def parse_ann_gene(value: Any) -> list[str]:
    text = clean(value)
    if norm_missing(text):
        return []
    genes: set[str] = set()
    for record in text.split(","):
        fields = record.split("|")
        for idx in (3, 4):
            if idx < len(fields):
                gene = norm_gene(fields[idx])
                if gene:
                    genes.add(gene)
    return sorted(genes)


def parse_csq_gene(value: Any) -> list[str]:
    text = clean(value)
    if norm_missing(text):
        return []
    genes: set[str] = set()
    for record in text.split(","):
        fields = record.split("|")
        for field in fields[:10]:
            gene = norm_gene(field)
            if gene and not gene.startswith("ENSG"):
                genes.add(gene)
    return sorted(genes)


def to_float(value: Any) -> float | None:
    text = clean(value)
    if norm_missing(text):
        return None
    text = text.rstrip("%")
    try:
        num = float(text)
    except ValueError:
        return None
    if "%" in clean(value):
        return num / 100
    return num


def valid_cancer_type(value: Any) -> str:
    text = clean(value)
    if norm_missing(text):
        return ""
    if to_float(text) is not None:
        return ""
    if len(text) > 80:
        return ""
    if "///" in text or text.count("/") > 2:
        return ""
    tokens = [tok for tok in re.split(r"[^A-Za-z0-9.+-]+", text) if tok]
    gene_like = sum(1 for tok in tokens if GENE_RE.match(tok.upper()))
    if len(tokens) >= 6 and gene_like >= max(4, len(tokens) // 2):
        return ""
    if COORD_RE.match(text):
        return ""
    if len(text) < 3:
        return ""
    return text


def normalize_cancer_type(value: Any, paper_id: str, paper_context: dict[str, dict[str, str]]) -> str:
    cancer = valid_cancer_type(value)
    if cancer:
        return cancer
    fallback = valid_cancer_type(paper_context.get(clean(paper_id), {}).get("cancer_type"))
    return fallback or "not_reported"


def load_paper_context(evidence_csv: Path) -> dict[str, dict[str, str]]:
    context: dict[str, dict[str, str]] = {}
    if not evidence_csv.exists():
        return context
    with evidence_csv.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            paper_id = clean(row.get("paper_id"))
            if not paper_id or paper_id in context:
                continue
            context[paper_id] = {
                "title": clean(row.get("title")),
                "year": clean(row.get("year")),
                "journal": clean(row.get("journal")),
                "doi": clean(row.get("doi")),
                "pmid": clean(row.get("pmid")),
                "cancer_type": clean(row.get("cancer_type")),
                "assay_type": clean(row.get("assay_type")),
                "sample_type": clean(row.get("sample_type")),
            }
    return context


def load_sources(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    rows = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            ext = clean(row.get("extension")).lower()
            stype = clean(row.get("source_document_type")).lower()
            if ext in {".xlsx", ".xls", ".csv", ".tsv"} and "supplementary" in stype:
                if clean(row.get("parser_status")).startswith("parsed"):
                    rows.append(row)
    return rows


def load_candidate_sheet_map(path: Path, min_candidate_rows: int, max_total: int = 0) -> dict[str, set[str]]:
    entries: list[tuple[int, str, str]] = []
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            try:
                candidate_rows = int(row.get("candidate_rows") or 0)
            except ValueError:
                candidate_rows = 0
            if candidate_rows >= min_candidate_rows:
                entries.append((candidate_rows, clean(row.get("source_document_id")), clean(row.get("source_sheet"))))
    entries.sort(reverse=True, key=lambda item: item[0])
    if max_total:
        entries = entries[:max_total]
    mapping: defaultdict[str, set[str]] = defaultdict(set)
    for _, source_id, sheet in entries:
        if source_id and sheet:
            mapping[source_id].add(sheet)
    return dict(mapping)


def load_candidate_chunk_map(path: Path, min_candidate_rows: int, max_total: int = 0) -> dict[str, dict[str, list[tuple[int, int]]]]:
    entries: list[tuple[int, str, str, int, int]] = []
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            try:
                candidate_rows = int(row.get("candidate_rows") or 0)
                row_start = int(row.get("row_start") or 1)
                row_end = int(row.get("row_end") or row_start)
            except ValueError:
                continue
            if candidate_rows >= min_candidate_rows:
                entries.append((candidate_rows, clean(row.get("source_document_id")), clean(row.get("source_sheet")), row_start, row_end))
    entries.sort(reverse=True, key=lambda item: item[0])
    if max_total:
        entries = entries[:max_total]
    mapping: defaultdict[str, defaultdict[str, list[tuple[int, int]]]] = defaultdict(lambda: defaultdict(list))
    for _, source_id, sheet, row_start, row_end in entries:
        if source_id and sheet:
            mapping[source_id][sheet].append((row_start, row_end))
    return {source_id: dict(sheets) for source_id, sheets in mapping.items()}


SAMPLE_HINTS = {
    "sample", "sample_id", "sampleid", "sample_name", "specimen", "specimen_id", "tumor_sample",
    "tumour_sample", "id", "variant_id", "sample_barcode"
}
PATIENT_HINTS = {"patient", "patient_id", "patient_no", "donor", "donor_id", "case", "case_id", "subject", "subject_id", "path_pid"}
GENE_HINTS = {"gene", "genes", "hugo_symbol", "hugo_gene_symbol", "gene_symbol", "symbol", "mutation_analysis_gene_variant", "mutated_gene"}
CHROM_HINTS = {"chrom", "chromosome", "chr", "#chrom", "chromo"}
POS_HINTS = {"pos", "position", "start", "chr_start", "genomic_position", "coordinate", "id"}
REF_HINTS = {"ref", "reference", "reference_allele", "reference_allele_variant", "ref_allele"}
ALT_HINTS = {"alt", "alternate", "alternate_allele", "altered_allele", "alt_allele", "tumor_seq_allele2"}
VAF_HINTS = {"vaf", "af", "allele_frequency", "variant_allele_frequency", "mutation_analysis_vaf_variant", "tumor_vaf"}
TYPE_HINTS = {"variant_classification", "variant_type", "type", "mutation_type", "mutation_effect", "consequence"}
ANN_HINTS = {"ann", "annotation"}
CSQ_HINTS = {"csq", "vep"}
TIME_HINTS = {"timepoint", "time_point", "baseline", "week", "cycle", "visit"}
CANCER_HINTS = {"cancer", "cancer_type", "histology", "histology_type", "tumor_type", "tumour_type", "diagnosis"}


def best_col(columns: list[str], hints: set[str], contains: tuple[str, ...] = ()) -> str:
    normed = {norm_header(col): col for col in columns}
    for hint in hints:
        if hint in normed:
            return normed[hint]
    for ncol, original in normed.items():
        if any(token in ncol for token in contains):
            return original
    return ""


def find_header_row(raw: pd.DataFrame) -> int | None:
    best_idx = None
    best_score = 0
    for idx in range(min(len(raw), 30)):
        values = [norm_header(v) for v in raw.iloc[idx].tolist()]
        joined = " ".join(values)
        score = 0
        if any(v in SAMPLE_HINTS or "sample" in v for v in values):
            score += 3
        if any(v in PATIENT_HINTS or "patient" in v or "donor" in v for v in values):
            score += 2
        if any(v in GENE_HINTS or v == "hugo" or "gene" in v for v in values):
            score += 3
        if any(v in CHROM_HINTS for v in values) and any(v in POS_HINTS for v in values):
            score += 3
        if "ann" in values or "csq" in values:
            score += 2
        if "variant" in joined or "mutation" in joined or "allele" in joined:
            score += 1
        non_empty = sum(1 for v in values if v)
        if non_empty < 3:
            score -= 2
        if score > best_score:
            best_score = score
            best_idx = idx
    return best_idx if best_score >= 4 else None


def dataframe_from_excel(path: Path, sheet_name: str) -> pd.DataFrame | None:
    try:
        raw = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=30, engine="openpyxl")
        header_idx = find_header_row(raw)
        if header_idx is None:
            return None
        df = pd.read_excel(path, sheet_name=sheet_name, header=header_idx, engine="openpyxl")
        return normalize_dataframe_columns(df)
    except Exception:
        return None


def unique_columns(values: list[Any]) -> list[str]:
    cols: list[str] = []
    seen: Counter[str] = Counter()
    for index, value in enumerate(values):
        base = clean(value) or f"unnamed_{index}"
        if base.lower().startswith("unnamed"):
            base = f"unnamed_{index}"
        seen[base] += 1
        cols.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return cols


def dataframe_from_excel_chunks(
    path: Path,
    sheet_name: str,
    chunks: list[tuple[int, int]],
    max_rows: int,
) -> pd.DataFrame | None:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            return None
        ws = wb[sheet_name]
        preview_values = []
        for row in ws.iter_rows(min_row=1, max_row=min(30, ws.max_row or 30), values_only=True):
            preview_values.append(list(row))
        header_idx = find_header_row(pd.DataFrame(preview_values))
        if header_idx is None:
            wb.close()
            return None
        header_row_number = header_idx + 1
        columns = unique_columns(preview_values[header_idx])
        width = len(columns)
        data_rows: list[list[Any]] = []
        seen_ranges: set[tuple[int, int]] = set()
        for row_start, row_end in sorted(chunks):
            start = max(row_start, header_row_number + 1)
            end = max(start, row_end)
            if (start, end) in seen_ranges:
                continue
            seen_ranges.add((start, end))
            for row in ws.iter_rows(min_row=start, max_row=end, values_only=True):
                values = list(row[:width])
                if any(not norm_missing(value) for value in values):
                    data_rows.append((values + [None] * width)[:width])
                if len(data_rows) >= max_rows:
                    break
            if len(data_rows) >= max_rows:
                break
        wb.close()
        if not data_rows:
            return None
        return pd.DataFrame(data_rows, columns=columns)
    except Exception:
        return None


def dataframe_from_text_table(path: Path, sep: str) -> pd.DataFrame | None:
    try:
        df = pd.read_csv(path, sep=sep, dtype=str, low_memory=False)
        return normalize_dataframe_columns(df)
    except Exception:
        return None


def normalize_dataframe_columns(df: pd.DataFrame) -> pd.DataFrame:
    cols = []
    seen: Counter[str] = Counter()
    for col in df.columns:
        base = clean(col) or "unnamed"
        if base.lower().startswith("unnamed"):
            base = f"unnamed_{len(cols)}"
        seen[base] += 1
        cols.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    df = df.copy()
    df.columns = cols
    return df


def parse_coord_from_id(value: Any) -> tuple[str, str, str, str]:
    text = clean(value)
    if norm_missing(text):
        return "", "", "", ""
    match = re.search(r"(?:chr)?([0-9]{1,2}|X|Y|M|MT)[:_](\d+)(?:[:_]?([ACGTN-]+)[/>_]([ACGTN-]+))?", text, re.I)
    if not match:
        return "", "", "", ""
    chrom = match.group(1).upper().replace("M", "MT")
    pos = match.group(2)
    ref = match.group(3) or ""
    alt = match.group(4) or ""
    return chrom, pos, ref, alt


def variant_key(chrom: str, pos: str, ref: str, alt: str, gene: str) -> str:
    if chrom and pos and ref and alt:
        return f"{chrom}:{pos}:{ref}>{alt}"
    if chrom and pos:
        return f"{chrom}:{pos}"
    return gene


def locus_feature_id(chrom: str, pos: str, ref: str, alt: str) -> str:
    if chrom and pos and ref and alt:
        return f"LOCUS:{normalize_chrom(chrom)}:{pos}:{ref}>{alt}"
    if chrom and pos:
        return f"LOCUS:{normalize_chrom(chrom)}:{pos}"
    return ""


def feature_from_event(gene: str, chrom: str, pos: str, ref: str, alt: str) -> dict[str, str]:
    if gene and gene != "not_annotated":
        return {"feature_id": f"GENE:{gene}", "feature_type": "gene", "feature_label": gene}
    locus = locus_feature_id(chrom, pos, ref, alt)
    if locus:
        return {"feature_id": locus, "feature_type": "locus", "feature_label": locus.replace("LOCUS:", "")}
    return {"feature_id": "UNANNOTATED:unknown", "feature_type": "unannotated", "feature_label": "unknown"}


def row_has_variant_signal(row: dict[str, Any], colmap: dict[str, str]) -> bool:
    if colmap.get("gene") and split_gene_values(row.get(colmap["gene"])):
        return True
    if colmap.get("ann") and parse_ann_gene(row.get(colmap["ann"])):
        return True
    if colmap.get("csq") and parse_csq_gene(row.get(colmap["csq"])):
        return True
    if colmap.get("chrom") and colmap.get("pos"):
        return bool(clean(row.get(colmap["chrom"])) and clean(row.get(colmap["pos"])))
    if colmap.get("pos"):
        return bool(parse_coord_from_id(row.get(colmap["pos"]))[0])
    return False


def infer_colmap(columns: list[str]) -> dict[str, str]:
    colmap = {
        "sample": best_col(columns, SAMPLE_HINTS, ("sample", "specimen", "barcode")),
        "patient": best_col(columns, PATIENT_HINTS, ("patient", "donor", "subject", "case")),
        "gene": best_col(columns, GENE_HINTS, ("gene", "hugo")),
        "chrom": best_col(columns, CHROM_HINTS, ("chrom",)),
        "pos": best_col(columns, POS_HINTS, ("position", "coord", "chr_start")),
        "ref": best_col(columns, REF_HINTS, ("ref", "reference")),
        "alt": best_col(columns, ALT_HINTS, ("alt", "altered", "alternate")),
        "vaf": best_col(columns, VAF_HINTS, ("vaf", "allele_frequency")),
        "variant_type": best_col(columns, TYPE_HINTS, ("variant_classification", "mutation_type", "consequence")),
        "ann": best_col(columns, ANN_HINTS, ("ann",)),
        "csq": best_col(columns, CSQ_HINTS, ("csq",)),
        "timepoint": best_col(columns, TIME_HINTS, ("timepoint", "baseline", "visit")),
        "cancer_type": best_col(columns, CANCER_HINTS, ("histology", "diagnosis", "cancer")),
    }
    return colmap


def table_is_sample_variant_matrix(df: pd.DataFrame, colmap: dict[str, str]) -> bool:
    columns = list(df.columns)
    has_sample = bool(colmap.get("sample") or colmap.get("patient"))
    has_gene_or_coord = bool(colmap.get("gene") or colmap.get("ann") or colmap.get("csq") or (colmap.get("chrom") and colmap.get("pos")))
    if not (has_sample and has_gene_or_coord):
        return False
    ncols = len(columns)
    if ncols > 120 and not (colmap.get("chrom") and colmap.get("pos")):
        return False
    return True


def extract_rows_from_table(
    df: pd.DataFrame,
    source: dict[str, str],
    sheet_name: str,
    paper_context: dict[str, dict[str, str]],
    max_rows: int,
    ncbi_index: dict[str, dict[str, dict[str, Any]]],
    coordinate_assemblies: list[str],
) -> list[dict[str, Any]]:
    colmap = infer_colmap(list(df.columns))
    if not table_is_sample_variant_matrix(df, colmap):
        return []

    paper_id = clean(source.get("paper_id"))
    context = paper_context.get(paper_id, {})
    rows: list[dict[str, Any]] = []
    for row_number, (_, record) in enumerate(df.head(max_rows).iterrows(), start=1):
        row = {col: record.get(col) for col in df.columns}
        if not row_has_variant_signal(row, colmap):
            continue

        sample_id = clean(row.get(colmap.get("sample", ""))) if colmap.get("sample") else ""
        patient_id = clean(row.get(colmap.get("patient", ""))) if colmap.get("patient") else ""
        if not sample_id and patient_id:
            sample_id = patient_id
        if not sample_id:
            continue

        genes: set[str] = set()
        for gene in split_gene_values(row.get(colmap.get("gene", "")) if colmap.get("gene") else ""):
            genes.add(gene)
        if colmap.get("ann"):
            genes.update(parse_ann_gene(row.get(colmap["ann"])))
        if colmap.get("csq"):
            genes.update(parse_csq_gene(row.get(colmap["csq"])))
        has_reported_gene = bool(genes)

        chrom = clean(row.get(colmap.get("chrom", ""))) if colmap.get("chrom") else ""
        pos = clean(row.get(colmap.get("pos", ""))) if colmap.get("pos") else ""
        ref = clean(row.get(colmap.get("ref", ""))) if colmap.get("ref") else ""
        alt = clean(row.get(colmap.get("alt", ""))) if colmap.get("alt") else ""
        if not chrom and pos:
            chrom, pos2, ref2, alt2 = parse_coord_from_id(pos)
            pos = pos2 or pos
            ref = ref or ref2
            alt = alt or alt2
        chrom = re.sub(r"^chr", "", chrom, flags=re.I).upper().replace("M", "MT") if chrom else ""

        if not genes and not (chrom and pos):
            continue
        coordinate_annotation = annotate_coordinate_gene(chrom, pos, ncbi_index, coordinate_assemblies)
        coordinate_genes = split_gene_values(coordinate_annotation.get("coordinate_gene"))
        if not genes and coordinate_genes:
            genes.update(coordinate_genes)
        if not genes:
            genes.add("not_annotated")

        vaf = to_float(row.get(colmap.get("vaf", ""))) if colmap.get("vaf") else None
        cancer_type = valid_cancer_type(row.get(colmap.get("cancer_type", ""))) if colmap.get("cancer_type") else ""
        if not cancer_type:
            cancer_type = valid_cancer_type(context.get("cancer_type"))
        if not cancer_type:
            cancer_type = "not_reported"

        for gene in sorted(genes):
            key = variant_key(chrom, pos, ref, alt, gene)
            feature = feature_from_event(gene, chrom, pos, ref, alt)
            if has_reported_gene and gene != "not_annotated":
                annotation_status = "gene_reported"
            elif gene != "not_annotated" and coordinate_annotation.get("coordinate_gene_status") == "ncbi_interval_match":
                annotation_status = "ncbi_coordinate_gene_rescue"
            else:
                annotation_status = "coordinate_only"
            if has_reported_gene and coordinate_genes:
                reported_gene_coordinate_status = (
                    "reported_gene_matches_ncbi_interval"
                    if gene in coordinate_genes
                    else "reported_gene_differs_from_ncbi_interval"
                )
            elif has_reported_gene:
                reported_gene_coordinate_status = coordinate_annotation.get("coordinate_gene_status", "")
            else:
                reported_gene_coordinate_status = ""
            event_id = hashlib.sha1(
                "|".join([paper_id, clean(source.get("source_document_id")), sheet_name, str(row_number), sample_id, gene, key]).encode("utf-8", "ignore")
            ).hexdigest()[:16]
            rows.append(
                {
                    "event_id": f"SUPPVAR-{event_id}",
                    "paper_id": paper_id,
                    "source_document_id": clean(source.get("source_document_id")),
                    "source_file": clean(source.get("source_file")),
                    "source_path": clean(source.get("source_path")),
                    "source_sheet": sheet_name,
                    "source_row": row_number,
                    "sample_id": sample_id,
                    "patient_id": patient_id,
                    "timepoint": clean(row.get(colmap.get("timepoint", ""))) if colmap.get("timepoint") else "",
                    "cancer_type": cancer_type,
                    "assay_type": clean(context.get("assay_type")) or "not_reported",
                    "sample_type": clean(context.get("sample_type")) or "not_reported",
                    "gene": gene,
                    "feature_id": feature["feature_id"],
                    "feature_type": feature["feature_type"],
                    "feature_label": feature["feature_label"],
                    "chrom": chrom,
                    "pos": pos,
                    "ref": ref,
                    "alt": alt,
                    "variant_key": key,
                    "variant_type": clean(row.get(colmap.get("variant_type", ""))) if colmap.get("variant_type") else "",
                    "vaf": "" if vaf is None else vaf,
                    "annotation_status": annotation_status,
                    "coordinate_gene": coordinate_annotation.get("coordinate_gene", ""),
                    "coordinate_gene_assembly": coordinate_annotation.get("coordinate_gene_assembly", ""),
                    "coordinate_gene_status": coordinate_annotation.get("coordinate_gene_status", ""),
                    "reported_gene_coordinate_status": reported_gene_coordinate_status,
                    "extraction_confidence": "high" if sample_id and gene != "not_annotated" else "medium",
                    "column_map_json": json.dumps(colmap, ensure_ascii=False, sort_keys=True),
                }
            )
    return rows


def iter_source_tables(
    source: dict[str, str],
    allowed_sheets: set[str] | None = None,
    chunk_sheets: dict[str, list[tuple[int, int]]] | None = None,
    max_sheets: int = 0,
    max_rows: int = 200000,
) -> list[tuple[str, pd.DataFrame]]:
    path = Path(clean(source.get("source_path")))
    ext = clean(source.get("extension")).lower()
    if not path.exists():
        return []
    tables: list[tuple[str, pd.DataFrame]] = []
    if ext in {".xlsx", ".xls"}:
        try:
            xls = pd.ExcelFile(path)
            scanned = 0
            for sheet in xls.sheet_names:
                if allowed_sheets is not None and sheet not in allowed_sheets:
                    continue
                if max_sheets and scanned >= max_sheets:
                    break
                scanned += 1
                if chunk_sheets and sheet in chunk_sheets:
                    df = dataframe_from_excel_chunks(path, sheet, chunk_sheets[sheet], max_rows)
                else:
                    df = dataframe_from_excel(path, sheet)
                if df is not None and len(df) > 0:
                    tables.append((sheet, df))
        except Exception:
            return []
    elif ext in {".csv", ".tsv"}:
        sep = "\t" if ext == ".tsv" else ","
        df = dataframe_from_text_table(path, sep)
        if df is not None and len(df) > 0:
            tables.append((path.stem, df))
    return tables


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def build_gene_matrix(variant_rows: list[dict[str, Any]], paper_context: dict[str, dict[str, str]] | None = None) -> list[dict[str, Any]]:
    paper_context = paper_context or {}
    seen: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
    support: defaultdict[tuple[str, str, str, str, str], set[str]] = defaultdict(set)
    for row in variant_rows:
        if row["gene"] == "not_annotated":
            continue
        cancer_type = normalize_cancer_type(row.get("cancer_type"), row.get("paper_id", ""), paper_context)
        timepoint = clean(row.get("timepoint"))
        key = (row["paper_id"], cancer_type, row["sample_id"], row["patient_id"], timepoint, row["gene"])
        support[key].add(row["event_id"])
        if key not in seen:
            seen[key] = {
                "sample_gene_id": "SUPPGENE-" + hashlib.sha1("|".join(key).encode()).hexdigest()[:16],
                "paper_id": row["paper_id"],
                "cancer_type": cancer_type,
                "sample_id": row["sample_id"],
                "patient_id": row["patient_id"],
                "timepoint": timepoint,
                "gene": row["gene"],
                "assay_type": row["assay_type"],
                "sample_type": row["sample_type"],
                "source_document_ids": set(),
                "variant_count": 0,
                "max_vaf": None,
                "event_ids": set(),
            }
        item = seen[key]
        item["source_document_ids"].add(row["source_document_id"])
        item["event_ids"].add(row["event_id"])
        item["variant_count"] += 1
        vaf = to_float(row.get("vaf"))
        if vaf is not None:
            item["max_vaf"] = vaf if item["max_vaf"] is None else max(item["max_vaf"], vaf)

    output = []
    for item in seen.values():
        output.append(
            {
                **{k: v for k, v in item.items() if k not in {"source_document_ids", "event_ids"}},
                "source_document_ids": ";".join(sorted(item["source_document_ids"])),
                "event_ids": ";".join(sorted(item["event_ids"])),
                "max_vaf": "" if item["max_vaf"] is None else item["max_vaf"],
            }
        )
    return sorted(output, key=lambda r: (r["cancer_type"], r["paper_id"], r["sample_id"], r.get("timepoint", ""), r["gene"]))


def build_feature_matrix(variant_rows: list[dict[str, Any]], paper_context: dict[str, dict[str, str]] | None = None) -> list[dict[str, Any]]:
    paper_context = paper_context or {}
    seen: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
    for row in variant_rows:
        feature_id = clean(row.get("feature_id"))
        if not feature_id:
            continue
        cancer_type = normalize_cancer_type(row.get("cancer_type"), row.get("paper_id", ""), paper_context)
        timepoint = clean(row.get("timepoint"))
        key = (row["paper_id"], cancer_type, row["sample_id"], row["patient_id"], timepoint, feature_id)
        if key not in seen:
            seen[key] = {
                "sample_feature_id": "SUPPFEAT-" + hashlib.sha1("|".join(key).encode()).hexdigest()[:16],
                "paper_id": row["paper_id"],
                "cancer_type": cancer_type,
                "sample_id": row["sample_id"],
                "patient_id": row["patient_id"],
                "timepoint": timepoint,
                "feature_id": feature_id,
                "feature_type": clean(row.get("feature_type")),
                "feature_label": clean(row.get("feature_label")),
                "gene": clean(row.get("gene")),
                "chrom": clean(row.get("chrom")),
                "pos": clean(row.get("pos")),
                "ref": clean(row.get("ref")),
                "alt": clean(row.get("alt")),
                "assay_type": row["assay_type"],
                "sample_type": row["sample_type"],
                "source_document_ids": set(),
                "variant_count": 0,
                "max_vaf": None,
                "event_ids": set(),
                "annotation_statuses": set(),
            }
        item = seen[key]
        item["source_document_ids"].add(row["source_document_id"])
        item["event_ids"].add(row["event_id"])
        item["annotation_statuses"].add(clean(row.get("annotation_status")))
        item["variant_count"] += 1
        vaf = to_float(row.get("vaf"))
        if vaf is not None:
            item["max_vaf"] = vaf if item["max_vaf"] is None else max(item["max_vaf"], vaf)

    output = []
    for item in seen.values():
        output.append(
            {
                **{k: v for k, v in item.items() if k not in {"source_document_ids", "event_ids", "annotation_statuses"}},
                "source_document_ids": ";".join(sorted(item["source_document_ids"])),
                "variant_count": item["variant_count"],
                "max_vaf": "" if item["max_vaf"] is None else item["max_vaf"],
                "event_ids": ";".join(sorted(item["event_ids"])),
                "annotation_statuses": ";".join(sorted(status for status in item["annotation_statuses"] if status)),
            }
        )
    return sorted(output, key=lambda r: (r["cancer_type"], r["paper_id"], r["sample_id"], r.get("timepoint", ""), r["feature_id"]))


def build_matrices_from_variant_csv(
    variant_csv: Path,
    gene_csv: Path,
    feature_csv: Path,
    paper_context: dict[str, dict[str, str]],
    valid_gene_symbols: set[str] | None = None,
) -> dict[str, int]:
    """Rebuild compact sample matrices from the already parsed event matrix."""
    gene_seen: dict[tuple[str, ...], dict[str, Any]] = {}
    feature_seen: dict[tuple[str, ...], dict[str, Any]] = {}
    variant_rows = 0
    with variant_csv.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            variant_rows += 1
            paper_id = clean(row.get("paper_id"))
            cancer_type = normalize_cancer_type(row.get("cancer_type"), paper_id, paper_context)
            sample_id = clean(row.get("sample_id"))
            patient_id = clean(row.get("patient_id"))
            timepoint = clean(row.get("timepoint"))
            assay_type = clean(row.get("assay_type")) or "not_reported"
            sample_type = clean(row.get("sample_type")) or "not_reported"
            event_id = clean(row.get("event_id"))
            source_document_id = clean(row.get("source_document_id"))
            raw_gene = clean(row.get("gene"))
            gene = "not_annotated" if raw_gene == "not_annotated" else norm_gene(raw_gene)
            if gene and gene != "not_annotated" and valid_gene_symbols and gene not in valid_gene_symbols:
                gene = ""
            chrom = clean(row.get("chrom"))
            pos = clean(row.get("pos"))
            ref = clean(row.get("ref"))
            alt = clean(row.get("alt"))
            vaf = to_float(row.get("vaf"))

            if gene and gene != "not_annotated":
                gene_key = (paper_id, cancer_type, sample_id, patient_id, timepoint, gene)
                if gene_key not in gene_seen:
                    gene_seen[gene_key] = {
                        "sample_gene_id": "SUPPGENE-" + hashlib.sha1("|".join(gene_key).encode()).hexdigest()[:16],
                        "paper_id": paper_id,
                        "cancer_type": cancer_type,
                        "sample_id": sample_id,
                        "patient_id": patient_id,
                        "timepoint": timepoint,
                        "gene": gene,
                        "assay_type": assay_type,
                        "sample_type": sample_type,
                        "source_document_ids": set(),
                        "variant_count": 0,
                        "max_vaf": None,
                        "event_ids": set(),
                    }
                gene_item = gene_seen[gene_key]
                if source_document_id:
                    gene_item["source_document_ids"].add(source_document_id)
                if event_id:
                    gene_item["event_ids"].add(event_id)
                gene_item["variant_count"] += 1
                if vaf is not None:
                    gene_item["max_vaf"] = vaf if gene_item["max_vaf"] is None else max(gene_item["max_vaf"], vaf)

            feature_id = clean(row.get("feature_id"))
            feature_type = clean(row.get("feature_type")) or "feature"
            feature_label = clean(row.get("feature_label")) or feature_id
            if feature_type == "gene":
                if gene and gene != "not_annotated":
                    feature_id = f"GENE:{gene}"
                    feature_label = gene
                else:
                    locus = locus_feature_id(chrom, pos, ref, alt)
                    if locus:
                        feature_id = locus
                        feature_type = "locus"
                        feature_label = locus.replace("LOCUS:", "")
                    else:
                        continue
            if not feature_id:
                continue
            feature_key = (paper_id, cancer_type, sample_id, patient_id, timepoint, feature_id)
            if feature_key not in feature_seen:
                feature_seen[feature_key] = {
                    "sample_feature_id": "SUPPFEAT-" + hashlib.sha1("|".join(feature_key).encode()).hexdigest()[:16],
                    "paper_id": paper_id,
                    "cancer_type": cancer_type,
                    "sample_id": sample_id,
                    "patient_id": patient_id,
                    "timepoint": timepoint,
                    "feature_id": feature_id,
                    "feature_type": feature_type,
                    "feature_label": feature_label,
                    "gene": gene,
                    "chrom": chrom,
                    "pos": pos,
                    "ref": ref,
                    "alt": alt,
                    "assay_type": assay_type,
                    "sample_type": sample_type,
                    "source_document_ids": set(),
                    "variant_count": 0,
                    "max_vaf": None,
                    "event_ids": set(),
                    "annotation_statuses": set(),
                }
            feature_item = feature_seen[feature_key]
            if source_document_id:
                feature_item["source_document_ids"].add(source_document_id)
            if event_id:
                feature_item["event_ids"].add(event_id)
            status = clean(row.get("annotation_status"))
            if status:
                feature_item["annotation_statuses"].add(status)
            feature_item["variant_count"] += 1
            if vaf is not None:
                feature_item["max_vaf"] = vaf if feature_item["max_vaf"] is None else max(feature_item["max_vaf"], vaf)
            if variant_rows % 250000 == 0:
                print(f"Rebuilt matrix state from {variant_rows} variant events ...")

    gene_rows = []
    for item in gene_seen.values():
        gene_rows.append(
            {
                **{k: v for k, v in item.items() if k not in {"source_document_ids", "event_ids"}},
                "source_document_ids": ";".join(sorted(item["source_document_ids"])),
                "event_ids": ";".join(sorted(item["event_ids"])),
                "max_vaf": "" if item["max_vaf"] is None else item["max_vaf"],
            }
        )
    feature_rows = []
    for item in feature_seen.values():
        feature_rows.append(
            {
                **{k: v for k, v in item.items() if k not in {"source_document_ids", "event_ids", "annotation_statuses"}},
                "source_document_ids": ";".join(sorted(item["source_document_ids"])),
                "event_ids": ";".join(sorted(item["event_ids"])),
                "annotation_statuses": ";".join(sorted(status for status in item["annotation_statuses"] if status)),
                "max_vaf": "" if item["max_vaf"] is None else item["max_vaf"],
            }
        )
    gene_fields = [
        "sample_gene_id", "paper_id", "cancer_type", "sample_id", "patient_id", "timepoint", "gene", "assay_type",
        "sample_type", "source_document_ids", "variant_count", "max_vaf", "event_ids",
    ]
    feature_fields = [
        "sample_feature_id", "paper_id", "cancer_type", "sample_id", "patient_id", "timepoint", "feature_id",
        "feature_type", "feature_label", "gene", "chrom", "pos", "ref", "alt", "assay_type", "sample_type",
        "source_document_ids", "variant_count", "max_vaf", "event_ids", "annotation_statuses",
    ]
    write_csv(gene_csv, sorted(gene_rows, key=lambda r: (r["cancer_type"], r["paper_id"], r["sample_id"], r.get("timepoint", ""), r["gene"])), gene_fields)
    write_csv(feature_csv, sorted(feature_rows, key=lambda r: (r["cancer_type"], r["paper_id"], r["sample_id"], r.get("timepoint", ""), r["feature_id"])), feature_fields)
    return {"variant_events": variant_rows, "sample_gene_rows": len(gene_rows), "sample_feature_rows": len(feature_rows)}


def sample_feature_sets(feature_rows: list[dict[str, Any]], paper_context: dict[str, dict[str, str]] | None = None) -> dict[tuple[str, str], dict[str, Any]]:
    paper_context = paper_context or {}
    samples: dict[tuple[str, str], dict[str, Any]] = {}
    for row in feature_rows:
        paper_id = clean(row.get("paper_id"))
        cancer_type = normalize_cancer_type(row.get("cancer_type"), paper_id, paper_context)
        sample_id = clean(row.get("sample_id"))
        patient_id = clean(row.get("patient_id"))
        timepoint = clean(row.get("timepoint"))
        assay_type = clean(row.get("assay_type"))
        key = (cancer_type, f"{paper_id}::{patient_id}::{sample_id}::{timepoint}::{assay_type}")
        if key not in samples:
            samples[key] = {
                "cancer_type": cancer_type,
                "sample_global_id": key[1],
                "paper_id": paper_id,
                "sample_id": sample_id,
                "patient_id": patient_id,
                "timepoint": timepoint,
                "assay_type": assay_type,
                "sample_type": clean(row.get("sample_type")),
                "features": set(),
                "feature_labels": {},
                "feature_types": {},
                "source_document_ids": set(),
                "event_ids_by_feature": defaultdict(set),
            }
        feature_id = row["feature_id"]
        samples[key]["features"].add(feature_id)
        samples[key]["feature_labels"][feature_id] = clean(row.get("feature_label")) or feature_id
        samples[key]["feature_types"][feature_id] = clean(row.get("feature_type")) or "feature"
        for source_id in clean(row.get("source_document_ids")).split(";"):
            if source_id:
                samples[key]["source_document_ids"].add(source_id)
        for event_id in clean(row.get("event_ids")).split(";"):
            if event_id and len(samples[key]["event_ids_by_feature"][feature_id]) < MAX_RULE_SUPPORT_EVENTS:
                samples[key]["event_ids_by_feature"][feature_id].add(event_id)
    return samples


def sample_gene_sets(gene_rows: list[dict[str, Any]], paper_context: dict[str, dict[str, str]] | None = None) -> dict[tuple[str, str], dict[str, Any]]:
    paper_context = paper_context or {}
    samples: dict[tuple[str, str], dict[str, Any]] = {}
    for row in gene_rows:
        paper_id = clean(row.get("paper_id"))
        cancer_type = normalize_cancer_type(row.get("cancer_type"), paper_id, paper_context)
        sample_id = clean(row.get("sample_id"))
        patient_id = clean(row.get("patient_id"))
        timepoint = clean(row.get("timepoint"))
        assay_type = clean(row.get("assay_type"))
        key = (cancer_type, f"{paper_id}::{patient_id}::{sample_id}::{timepoint}::{assay_type}")
        if key not in samples:
            samples[key] = {
                "cancer_type": cancer_type,
                "sample_global_id": key[1],
                "paper_id": paper_id,
                "sample_id": sample_id,
                "patient_id": patient_id,
                "timepoint": timepoint,
                "assay_type": assay_type,
                "genes": set(),
                "source_document_ids": set(),
                "event_ids_by_gene": defaultdict(set),
            }
        samples[key]["genes"].add(row["gene"])
        for source_id in clean(row.get("source_document_ids")).split(";"):
            if source_id:
                samples[key]["source_document_ids"].add(source_id)
        for event_id in clean(row.get("event_ids")).split(";"):
            if event_id:
                samples[key]["event_ids_by_gene"][row["gene"]].add(event_id)
    return samples


def paper_metadata_for_ids(paper_ids: set[str], paper_context: dict[str, dict[str, str]]) -> dict[str, str]:
    ids = sorted(pid for pid in paper_ids if pid)
    journals = sorted({paper_context.get(pid, {}).get("journal", "") for pid in ids if paper_context.get(pid, {}).get("journal")})
    titles = [paper_context.get(pid, {}).get("title", "") for pid in ids if paper_context.get(pid, {}).get("title")]
    dois = sorted({paper_context.get(pid, {}).get("doi", "") for pid in ids if paper_context.get(pid, {}).get("doi")})
    pmids = sorted({paper_context.get(pid, {}).get("pmid", "") for pid in ids if paper_context.get(pid, {}).get("pmid")})
    return {
        "supporting_paper_ids": ";".join(ids),
        "supporting_journals": ";".join(journals),
        "supporting_titles": " | ".join(titles[:5]),
        "supporting_dois": ";".join(dois),
        "supporting_pmids": ";".join(pmids),
    }


def build_baselines(feature_rows: list[dict[str, Any]], min_samples: int, paper_context: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    samples = sample_feature_sets(feature_rows, paper_context)
    cancer_samples: defaultdict[str, set[str]] = defaultdict(set)
    cancer_feature_samples: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    cancer_feature_papers: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    cancer_feature_sources: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    cancer_feature_events: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    feature_labels: dict[tuple[str, str], str] = {}
    feature_types: dict[tuple[str, str], str] = {}
    for (_, sample_id), item in samples.items():
        cancer = item["cancer_type"]
        cancer_samples[cancer].add(sample_id)
        for feature_id in item["features"]:
            key = (cancer, feature_id)
            cancer_feature_samples[key].add(sample_id)
            cancer_feature_papers[key].add(item["paper_id"])
            cancer_feature_sources[key].update(item["source_document_ids"])
            cancer_feature_events[key].update(item["event_ids_by_feature"].get(feature_id, set()))
            feature_labels[key] = item["feature_labels"].get(feature_id, feature_id)
            feature_types[key] = item["feature_types"].get(feature_id, "feature")
    rows = []
    for (cancer, feature_id), sample_ids in cancer_feature_samples.items():
        n = len(cancer_samples[cancer])
        if n < min_samples:
            continue
        count = len(sample_ids)
        meta = paper_metadata_for_ids(cancer_feature_papers[(cancer, feature_id)], paper_context)
        feature_label = feature_labels.get((cancer, feature_id), feature_id)
        feature_type = feature_types.get((cancer, feature_id), "feature")
        rows.append(
            {
                "baseline_id": "BASE-" + hashlib.sha1(f"{cancer}|{feature_id}".encode()).hexdigest()[:16],
                "cancer_type": cancer,
                "gene": feature_label if feature_type == "gene" else "",
                "feature_id": feature_id,
                "feature_type": feature_type,
                "feature_label": feature_label,
                "sample_count": n,
                "gene_sample_count": count,
                "gene_frequency": round(count / n, 6) if n else 0,
                "rule_status": "cohort_supported" if count >= 3 else "rare_observation",
                "source_layer": "supplementary_material_cohort",
                **meta,
                "supporting_source_document_ids": ";".join(sorted(cancer_feature_sources[(cancer, feature_id)])),
                "supporting_event_ids": ";".join(sorted(cancer_feature_events[(cancer, feature_id)])[:50]),
            }
        )
    return sorted(rows, key=lambda r: (r["cancer_type"], -float(r["gene_frequency"]), r["feature_id"]))


def build_pair_rules(
    feature_rows: list[dict[str, Any]],
    min_samples: int,
    min_pair_support: int,
    min_lift: float,
    max_mutex_or: float,
    mutex_max_p_value: float,
    min_mutex_feature_count: int,
    min_mutex_expected: float,
    mutex_feature_types: set[str],
    max_features_per_sample_for_pairs: int,
    skip_not_reported_pair_rules: bool,
    paper_context: dict[str, dict[str, str]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    samples = sample_feature_sets(feature_rows, paper_context)
    by_cancer: defaultdict[str, list[set[str]]] = defaultdict(list)
    by_cancer_items: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in samples.values():
        by_cancer[item["cancer_type"]].append(set(item["features"]))
        by_cancer_items[item["cancer_type"]].append(item)

    co_rows: list[dict[str, Any]] = []
    mutex_rows: list[dict[str, Any]] = []
    stats: dict[str, Any] = {
        "sample_units_total": len(samples),
        "sample_units_pair_eligible": 0,
        "sample_units_skipped_high_feature": 0,
        "sample_units_skipped_not_reported_cancer": 0,
        "cancers_considered": 0,
        "cancers_with_rules": 0,
    }

    for cancer in by_cancer:
        mutex_count_before = len(mutex_rows)
        if skip_not_reported_pair_rules and cancer == "not_reported":
            stats["sample_units_skipped_not_reported_cancer"] += len(by_cancer_items[cancer])
            continue
        cancer_items_all = by_cancer_items[cancer]
        cancer_items = []
        for item in cancer_items_all:
            if max_features_per_sample_for_pairs and len(item["features"]) > max_features_per_sample_for_pairs:
                stats["sample_units_skipped_high_feature"] += 1
                continue
            cancer_items.append(item)
        stats["sample_units_pair_eligible"] += len(cancer_items)
        n = len(cancer_items)
        if n < min_samples:
            continue
        stats["cancers_considered"] += 1
        feature_counts: Counter[str] = Counter()
        pair_counts: Counter[tuple[str, str]] = Counter()
        all_features: set[str] = set()
        feature_papers: defaultdict[str, set[str]] = defaultdict(set)
        feature_sources: defaultdict[str, set[str]] = defaultdict(set)
        feature_events: defaultdict[str, set[str]] = defaultdict(set)
        feature_labels: dict[str, str] = {}
        feature_types: dict[str, str] = {}

        for item in cancer_items:
            features = item["features"]
            all_features.update(features)
            for feature in features:
                feature_counts[feature] += 1
                add_limited(feature_papers[feature], [item["paper_id"]], MAX_RULE_SUPPORT_SOURCES)
                add_limited(feature_sources[feature], item["source_document_ids"], MAX_RULE_SUPPORT_SOURCES)
                add_limited(feature_events[feature], item["event_ids_by_feature"].get(feature, set()), MAX_RULE_SUPPORT_EVENTS)
                feature_labels[feature] = item["feature_labels"].get(feature, feature)
                feature_types[feature] = item["feature_types"].get(feature, "feature")
            for a, b in combinations(sorted(features), 2):
                pair_counts[(a, b)] += 1

        common_mutex_features = [
            feature
            for feature, count in feature_counts.items()
            if count >= max(min_pair_support, min_mutex_feature_count)
            and (not mutex_feature_types or feature_types.get(feature, "feature") in mutex_feature_types)
        ]

        def pair_base(feature_a: str, feature_b: str, n_ab: int, odds_ratio: float, p_value: float, lift: float) -> dict[str, Any]:
            n_a = feature_counts[feature_a]
            n_b = feature_counts[feature_b]
            label_a = feature_labels.get(feature_a, feature_a)
            label_b = feature_labels.get(feature_b, feature_b)
            type_a = feature_types.get(feature_a, "feature")
            type_b = feature_types.get(feature_b, "feature")
            support_papers = feature_papers[feature_a] | feature_papers[feature_b]
            meta = paper_metadata_for_ids(support_papers, paper_context)
            return {
                "cancer_type": cancer,
                "gene_a": label_a if type_a == "gene" else "",
                "gene_b": label_b if type_b == "gene" else "",
                "feature_a_id": feature_a,
                "feature_b_id": feature_b,
                "feature_a_type": type_a,
                "feature_b_type": type_b,
                "feature_a_label": label_a,
                "feature_b_label": label_b,
                "sample_count": n,
                "gene_a_count": n_a,
                "gene_b_count": n_b,
                "pair_count": n_ab,
                "gene_a_frequency": round(n_a / n, 6),
                "gene_b_frequency": round(n_b / n, 6),
                "pair_frequency": round(n_ab / n, 6),
                "lift": round(lift, 6),
                "odds_ratio": "inf" if math.isinf(odds_ratio) else round(odds_ratio, 6),
                "fisher_p_value": round(p_value, 8),
                "source_layer": "supplementary_material_cohort",
                **meta,
                "supporting_source_document_ids": ";".join(sorted((feature_sources[feature_a] | feature_sources[feature_b]))),
                "supporting_event_ids": ";".join(sorted((feature_events[feature_a] | feature_events[feature_b]))[:80]),
            }

        co_candidate_pairs: set[tuple[str, str]] = set()
        co_pair_stats: dict[tuple[str, str], tuple[float, float, float]] = {}
        for feature_a, feature_b in pair_counts:
            n_a = feature_counts[feature_a]
            n_b = feature_counts[feature_b]
            n_ab = pair_counts.get((feature_a, feature_b), 0)
            if n_a < min_pair_support or n_b < min_pair_support:
                continue
            a_only = n_a - n_ab
            b_only = n_b - n_ab
            neither = n - n_ab - a_only - b_only
            table = [[n_ab, a_only], [b_only, neither]]
            odds_ratio, p_value = fisher_exact(table)
            expected = (n_a / n) * (n_b / n) * n
            lift = (n_ab / expected) if expected > 0 else 0
            if n_ab >= min_pair_support and lift >= min_lift:
                co_candidate_pairs.add((feature_a, feature_b))
                co_pair_stats[(feature_a, feature_b)] = (odds_ratio, p_value, lift)
            if n_ab > 0 and lift < 0.5:
                or_num = odds_ratio if not math.isinf(odds_ratio) else 999999
                type_a = feature_types.get(feature_a, "feature")
                type_b = feature_types.get(feature_b, "feature")
                if (
                    or_num <= max_mutex_or
                    and p_value <= mutex_max_p_value
                    and (not mutex_feature_types or (type_a in mutex_feature_types and type_b in mutex_feature_types))
                ):
                    base = pair_base(feature_a, feature_b, n_ab, odds_ratio, p_value, lift)
                    mutex_rows.append(
                        {
                            "rule_id": "MUTEX-" + hashlib.sha1(f"{cancer}|{feature_a}|{feature_b}".encode()).hexdigest()[:16],
                            **base,
                            "rule_type": "gene_mutual_exclusion",
                            "rule_status": "candidate_rule",
                        }
                    )

        pair_papers: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
        pair_sources: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
        pair_events: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
        if co_candidate_pairs:
            candidate_features = {feature for pair in co_candidate_pairs for feature in pair}
            for item in cancer_items:
                features = sorted(item["features"] & candidate_features)
                for a, b in combinations(features, 2):
                    pair = (a, b)
                    if pair not in co_candidate_pairs:
                        continue
                    add_limited(pair_papers[pair], [item["paper_id"]], MAX_RULE_SUPPORT_SOURCES)
                    add_limited(pair_sources[pair], item["source_document_ids"], MAX_RULE_SUPPORT_SOURCES)
                    add_limited(pair_events[pair], item["event_ids_by_feature"].get(a, set()), MAX_RULE_SUPPORT_EVENTS)
                    add_limited(pair_events[pair], item["event_ids_by_feature"].get(b, set()), MAX_RULE_SUPPORT_EVENTS)

        for feature_a, feature_b in co_candidate_pairs:
            odds_ratio, p_value, lift = co_pair_stats[(feature_a, feature_b)]
            n_ab = pair_counts[(feature_a, feature_b)]
            base = pair_base(feature_a, feature_b, n_ab, odds_ratio, p_value, lift)
            if pair_papers.get((feature_a, feature_b)):
                meta = paper_metadata_for_ids(pair_papers[(feature_a, feature_b)], paper_context)
                base.update(meta)
            if pair_sources.get((feature_a, feature_b)):
                base["supporting_source_document_ids"] = ";".join(sorted(pair_sources[(feature_a, feature_b)]))
            if pair_events.get((feature_a, feature_b)):
                base["supporting_event_ids"] = ";".join(sorted(pair_events[(feature_a, feature_b)])[:80])
            status = "high_confidence_rule" if p_value < 0.05 and n_ab >= 5 else "candidate_rule"
            co_rows.append(
                {
                    "rule_id": "COOCCUR-" + hashlib.sha1(f"{cancer}|{feature_a}|{feature_b}".encode()).hexdigest()[:16],
                    **base,
                    "rule_type": "gene_cooccurrence",
                    "rule_status": status,
                }
            )

        for feature_a, feature_b in combinations(sorted(common_mutex_features), 2):
            if (feature_a, feature_b) in pair_counts:
                continue
            n_a = feature_counts[feature_a]
            n_b = feature_counts[feature_b]
            expected = (n_a / n) * (n_b / n) * n
            if expected < min_mutex_expected:
                continue
            n_ab = 0
            a_only = n_a
            b_only = n_b
            neither = n - n_a - n_b
            if neither < 0:
                continue
            odds_ratio, p_value = fisher_exact([[n_ab, a_only], [b_only, neither]])
            if p_value > mutex_max_p_value:
                continue
            base = pair_base(feature_a, feature_b, n_ab, odds_ratio, p_value, 0)
            mutex_rows.append(
                {
                    "rule_id": "MUTEX-" + hashlib.sha1(f"{cancer}|{feature_a}|{feature_b}".encode()).hexdigest()[:16],
                    **base,
                    "rule_type": "gene_mutual_exclusion",
                    "rule_status": "candidate_rule",
                }
            )

        if co_candidate_pairs or len(mutex_rows) > mutex_count_before:
            stats["cancers_with_rules"] += 1
    co_rows.sort(key=lambda r: (r["cancer_type"], -r["pair_count"], -float(r["lift"])))
    mutex_rows.sort(key=lambda r: (r["cancer_type"], r["pair_count"], float(r["fisher_p_value"])))
    stats["cooccurrence_rules"] = len(co_rows)
    stats["mutex_rules"] = len(mutex_rows)
    return co_rows, mutex_rows, stats


def write_report(path: Path, stats: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Supplementary Cohort Rule Build Report",
        "",
        "## Summary",
        "",
        f"- Sources scanned: {stats['sources_scanned']}",
        f"- Tables accepted: {stats['tables_accepted']}",
        f"- Variant events: {stats['variant_events']}",
        f"- Sample-gene rows: {stats['sample_gene_rows']}",
        f"- Sample-feature rows: {stats['sample_feature_rows']}",
        f"- Cancer-gene baselines: {stats['baseline_rows']}",
        f"- Co-occurrence rules: {stats['cooccurrence_rules']}",
        f"- Mutual-exclusion candidate rules: {stats['mutex_rules']}",
    ]
    pair_stats = stats.get("pair_rule_stats")
    if isinstance(pair_stats, dict):
        lines.extend(
            [
                f"- Pair-rule sample units total: {pair_stats.get('sample_units_total', '')}",
                f"- Pair-rule sample units eligible: {pair_stats.get('sample_units_pair_eligible', '')}",
                f"- Pair-rule sample units skipped for high feature load: {pair_stats.get('sample_units_skipped_high_feature', '')}",
                f"- Pair-rule sample units skipped for missing cancer context: {pair_stats.get('sample_units_skipped_not_reported_cancer', '')}",
            ]
        )
    lines.extend(["", "## Top Accepted Tables", ""])
    for item in stats["top_tables"][:20]:
        lines.append(f"- {item['event_count']} rows | {item['paper_id']} | {item['source_sheet']} | `{item['source_file']}`")
    lines.extend(
        [
            "",
            "## Notes",
            "",
            "- This is a supplementary-material cohort/statistical layer, not a replacement for the literature evidence KG.",
            "- Rules are sample-level only when sample or patient identifiers are present in the source table.",
            "- Coordinate-only rows are retained as locus-level features; when NCBI interval annotation can rescue a gene, the rescued gene and assembly are preserved.",
            "- VAF values are preserved from source tables; downstream QC should normalize percent/fraction units and flag values outside 0-1.",
            "- Excel workbooks can be scanned with the same script, but very large workbooks should be processed in bounded batches.",
            "- Candidate rules require review before automated clinical or MRD QC use.",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    args = parse_args()
    coordinate_assemblies = [item.strip() for item in args.coordinate_assemblies.split(",") if item.strip()]
    ncbi_index = load_ncbi_index(
        {
            "GRCh37": artifact_path(args.ncbi_grch37_json),
            "GRCh38": artifact_path(args.ncbi_grch38_json),
        }
    )
    valid_gene_symbols = ncbi_gene_symbols_from_index(ncbi_index)
    paper_context = load_paper_context(artifact_path(args.evidence_csv))
    provenance_fields = [
        "supporting_paper_ids", "supporting_journals", "supporting_titles", "supporting_dois",
        "supporting_pmids", "supporting_source_document_ids", "supporting_event_ids",
    ]
    baseline_fields = [
        "baseline_id", "cancer_type", "gene", "feature_id", "feature_type", "feature_label", "sample_count", "gene_sample_count",
        "gene_frequency", "rule_status", "source_layer", *provenance_fields,
    ]
    rule_fields = [
        "rule_id", "rule_type", "rule_status", "cancer_type", "gene_a", "gene_b", "sample_count", "gene_a_count",
        "gene_b_count", "pair_count", "gene_a_frequency", "gene_b_frequency", "pair_frequency", "lift", "odds_ratio",
        "fisher_p_value", "source_layer", "feature_a_id", "feature_b_id", "feature_a_type", "feature_b_type",
        "feature_a_label", "feature_b_label", *provenance_fields,
    ]
    mutex_feature_types = {item.strip() for item in args.mutex_feature_types.split(",") if item.strip()}
    pair_rule_args = (
        args.min_samples,
        args.min_pair_support,
        args.min_lift,
        args.max_mutex_odds_ratio,
        args.mutex_max_p_value,
        args.min_mutex_feature_count,
        args.min_mutex_expected,
        mutex_feature_types,
        args.max_features_per_sample_for_pairs,
        args.skip_not_reported_pair_rules,
        paper_context,
    )
    if args.from_variant_csv:
        matrix_stats = build_matrices_from_variant_csv(
            artifact_path(args.from_variant_csv),
            artifact_path(args.out_gene_csv),
            artifact_path(args.out_feature_csv),
            paper_context,
            valid_gene_symbols,
        )
        feature_rows = read_csv_rows(artifact_path(args.out_feature_csv))
        baseline_rows = build_baselines(feature_rows, args.min_samples, paper_context)
        write_csv(artifact_path(args.out_baseline_csv), baseline_rows, baseline_fields)
        co_rows, mutex_rows, pair_stats = build_pair_rules(feature_rows, *pair_rule_args)
        write_csv(artifact_path(args.out_cooccur_csv), co_rows, rule_fields)
        write_csv(artifact_path(args.out_mutex_csv), mutex_rows, rule_fields)
        stats = {
            "sources_scanned": 0,
            "tables_accepted": 0,
            **matrix_stats,
            "baseline_rows": len(baseline_rows),
            "cooccurrence_rules": len(co_rows),
            "mutex_rules": len(mutex_rows),
            "pair_rule_stats": pair_stats,
            "top_tables": [],
        }
        write_report(artifact_path(args.out_report_md), stats)
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        return
    if args.from_feature_csv:
        feature_rows = read_csv_rows(artifact_path(args.from_feature_csv))
        baseline_rows = build_baselines(feature_rows, args.min_samples, paper_context)
        write_csv(artifact_path(args.out_baseline_csv), baseline_rows, baseline_fields)
        co_rows, mutex_rows, pair_stats = build_pair_rules(feature_rows, *pair_rule_args)
        write_csv(artifact_path(args.out_cooccur_csv), co_rows, rule_fields)
        write_csv(artifact_path(args.out_mutex_csv), mutex_rows, rule_fields)
        stats = {
            "sources_scanned": 0,
            "tables_accepted": 0,
            "variant_events": "not_recomputed",
            "sample_gene_rows": "not_recomputed",
            "sample_feature_rows": len(feature_rows),
            "baseline_rows": len(baseline_rows),
            "cooccurrence_rules": len(co_rows),
            "mutex_rules": len(mutex_rows),
            "pair_rule_stats": pair_stats,
            "top_tables": [],
        }
        write_report(artifact_path(args.out_report_md), stats)
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        return

    sources = load_sources(artifact_path(args.sources_csv))
    allowed_extensions = {item.strip().lower() for item in args.extensions.split(",") if item.strip()}
    if allowed_extensions:
        sources = [source for source in sources if clean(source.get("extension")).lower() in allowed_extensions]
    candidate_sheet_map = load_candidate_sheet_map(artifact_path(args.excel_chunks_csv), args.min_candidate_rows, args.max_candidate_sheets_total)
    candidate_chunk_map = load_candidate_chunk_map(artifact_path(args.excel_chunks_csv), args.min_candidate_rows, args.max_candidate_sheets_total)
    if args.candidate_sheets_only:
        filtered = []
        excel_seen = 0
        for source in sources:
            ext = clean(source.get("extension")).lower()
            source_id = clean(source.get("source_document_id"))
            if args.max_source_size_mb:
                try:
                    size_mb = int(source.get("size_bytes") or 0) / (1024 * 1024)
                except ValueError:
                    size_mb = 0
                if size_mb > args.max_source_size_mb:
                    continue
            if ext in {".xlsx", ".xls"} and source_id not in candidate_sheet_map:
                continue
            if ext in {".xlsx", ".xls"}:
                excel_seen += 1
                if args.max_excel_sources and excel_seen > args.max_excel_sources:
                    continue
            filtered.append(source)
        sources = filtered
    if args.max_files:
        sources = sources[: args.max_files]
    variant_rows: list[dict[str, Any]] = []
    accepted_tables: list[dict[str, Any]] = []
    for idx, source in enumerate(sources, 1):
        if idx % 50 == 0:
            print(f"Scanning supplementary source {idx}/{len(sources)} ...")
        allowed_sheets = candidate_sheet_map.get(clean(source.get("source_document_id"))) if args.candidate_sheets_only else None
        chunk_sheets = candidate_chunk_map.get(clean(source.get("source_document_id"))) if args.candidate_sheets_only else None
        for sheet_name, df in iter_source_tables(source, allowed_sheets, chunk_sheets, args.max_sheets_per_source, args.max_rows_per_sheet):
            rows = extract_rows_from_table(df, source, sheet_name, paper_context, args.max_rows_per_sheet, ncbi_index, coordinate_assemblies)
            if rows:
                accepted_tables.append(
                    {
                        "paper_id": clean(source.get("paper_id")),
                        "source_document_id": clean(source.get("source_document_id")),
                        "source_file": clean(source.get("source_file")),
                        "source_sheet": sheet_name,
                        "event_count": len(rows),
                    }
                )
                variant_rows.extend(rows)

    variant_fields = [
        "event_id", "paper_id", "source_document_id", "source_file", "source_path", "source_sheet", "source_row",
        "sample_id", "patient_id", "timepoint", "cancer_type", "assay_type", "sample_type", "gene", "feature_id",
        "feature_type", "feature_label", "chrom", "pos", "ref", "alt", "variant_key", "variant_type", "vaf",
        "annotation_status", "coordinate_gene", "coordinate_gene_assembly", "coordinate_gene_status",
        "reported_gene_coordinate_status", "extraction_confidence", "column_map_json",
    ]
    write_csv(artifact_path(args.out_variant_csv), variant_rows, variant_fields)

    gene_rows = build_gene_matrix(variant_rows, paper_context)
    gene_fields = [
        "sample_gene_id", "paper_id", "cancer_type", "sample_id", "patient_id", "timepoint", "gene", "assay_type",
        "sample_type", "source_document_ids", "variant_count", "max_vaf", "event_ids",
    ]
    write_csv(artifact_path(args.out_gene_csv), gene_rows, gene_fields)

    feature_rows = build_feature_matrix(variant_rows, paper_context)
    feature_fields = [
        "sample_feature_id", "paper_id", "cancer_type", "sample_id", "patient_id", "timepoint", "feature_id",
        "feature_type", "feature_label", "gene", "chrom", "pos", "ref", "alt", "assay_type", "sample_type",
        "source_document_ids", "variant_count", "max_vaf", "event_ids", "annotation_statuses",
    ]
    write_csv(artifact_path(args.out_feature_csv), feature_rows, feature_fields)

    baseline_rows = build_baselines(feature_rows, args.min_samples, paper_context)
    write_csv(artifact_path(args.out_baseline_csv), baseline_rows, baseline_fields)

    co_rows, mutex_rows, pair_stats = build_pair_rules(feature_rows, *pair_rule_args)
    write_csv(artifact_path(args.out_cooccur_csv), co_rows, rule_fields)
    write_csv(artifact_path(args.out_mutex_csv), mutex_rows, rule_fields)

    accepted_tables.sort(key=lambda item: item["event_count"], reverse=True)
    stats = {
        "sources_scanned": len(sources),
        "tables_accepted": len(accepted_tables),
        "variant_events": len(variant_rows),
        "sample_gene_rows": len(gene_rows),
        "sample_feature_rows": len(feature_rows),
        "baseline_rows": len(baseline_rows),
        "cooccurrence_rules": len(co_rows),
        "mutex_rules": len(mutex_rows),
        "pair_rule_stats": pair_stats,
        "top_tables": accepted_tables[:30],
    }
    write_report(artifact_path(args.out_report_md), stats)
    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
